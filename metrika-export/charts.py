"""Панель с диаграммами: графики на seaborn → лист «Панель» в Excel и dashboard.html.

Правила оформления — сдержанные: один цвет на рейтинг (у источников и страниц
нет порядка, поэтому «чем больше, тем темнее» здесь вредно), прошлый период —
серым рядом с текущим, тонкие сетки, никаких двух осей на одном графике.
Изменения окрашены по смыслу: рост отказов — плохо, хотя число растёт.
"""
import base64
from pathlib import Path

# Цвета: акцент — синий, прошлый период — серый, «лучше / хуже» — зелёный / красный
ACCENT = '#2a78d6'
PREVIOUS = '#b9b8b1'
GOOD = '#0ca30c'
BAD = '#d03b3b'
INK = '#0b0b0b'
INK_2 = '#52514e'
MUTED = '#898781'
GRID = '#e1e0d9'
BASELINE = '#c3c2b7'
SURFACE = '#fcfcfb'
GOAL_C = '#ee7b2f'      # достижения целей — оранжевым рядом с синими визитами

# Шапка панели и страниц отчётов — как на sait-activator.ru: логотип (logo.png рядом со
# скриптом) и две кнопки. Свои ссылки и подписи — здесь
BRAND = {
    'site': 'https://sait-activator.ru/',
    'logo_alt': 'Сайт Активатор — агентство интернет-маркетинга',
    'order': ('Заказать SEO и GEO-продвижение', 'https://sait-activator.ru/'),
    'telegram': ('Подписаться на Telegram-канал', 'https://t.me/seo_python_2neuron'),
}
TG_ICON = ('<svg viewBox="0 0 24 24" width="18" height="18" aria-hidden="true"><path fill="currentColor" '
           'd="M9.04 15.47 8.67 20.7c.53 0 .76-.23 1.04-.5l2.5-2.39 5.18 3.79c.95.52 1.62.25 1.88-.88l3.4-15.95'
           'c.31-1.43-.52-1.99-1.45-1.64L1.2 10.9c-1.38.54-1.36 1.31-.24 1.66l5.1 1.59L17.9 6.7c.56-.37 1.07-.17.65.2'
           'z"/></svg>')
BRAND_CSS = """
.brand{background:#fff;border-bottom:1px solid #e1e0d9}
.brand-in{margin:0 auto;padding:14px 16px;display:flex;align-items:center;justify-content:space-between;gap:16px;
  flex-wrap:wrap}
.brand .logo img{height:46px;width:auto;display:block}
.brand .logo b{font-size:22px;color:#001f34;letter-spacing:.5px}
.cta{display:flex;gap:10px;flex-wrap:wrap}
.btn{display:inline-flex;align-items:center;justify-content:center;gap:8px;padding:11px 20px;border-radius:4px;
  font:600 14.5px system-ui,-apple-system,"Segoe UI",sans-serif;text-decoration:none;white-space:nowrap}
.btn:hover{text-decoration:none}
.btn.order{background:#ff7729;color:#fff;border:1.5px solid #ff7729}
.btn.order:hover{background:#f0651a;border-color:#f0651a}
.btn.tg{background:#fff;color:#001f34;border:1.5px solid #5cc9fa} .btn.tg svg{color:#229ed9}
.btn.tg:hover{background:#f1faff}
@media (max-width:640px){.brand-in{justify-content:center}.cta{width:100%}.btn{flex:1 1 100%}}
"""


def brand_header(width=1100):
    """Белая полоса с логотипом и кнопками; логотип встроен в страницу — интернет не нужен."""
    import html as _html
    logo_file = Path(__file__).resolve().parent / 'logo.png'
    alt = _html.escape(BRAND['logo_alt'])
    logo = (f'<img src="data:image/png;base64,{base64.b64encode(logo_file.read_bytes()).decode()}" alt="{alt}">'
            if logo_file.is_file() else f'<b>{alt}</b>')
    (order_text, order_url), (tg_text, tg_url) = BRAND['order'], BRAND['telegram']
    return (f'<header class="brand"><div class="brand-in" style="max-width:{width}px">'
            f'<a class="logo" href="{BRAND["site"]}" target="_blank" rel="noopener">{logo}</a>'
            f'<div class="cta"><a class="btn tg" href="{tg_url}" target="_blank" rel="noopener">{TG_ICON}'
            f'{_html.escape(tg_text)}</a><a class="btn order" href="{order_url}" target="_blank" rel="noopener">'
            f'{_html.escape(order_text)}</a></div></div></header>')


# Подписи при выбранной главной цели: «Достижения целей» → «Достижения цели» и т. п.
RELABEL = {}


def lbl(label):
    return RELABEL.get(label, label)


# Для каких показателей рост — это хорошо
UP_IS_GOOD = {
    'Визиты': True, 'Посетители': True, 'Просмотры': True, 'Отказы, %': False,
    'Глубина просмотра': True, 'Время на сайте': True, 'Новые посетители, %': None,
    'Достижения целей': True, 'Конверсия в любую цель, %': True,
}


def available():
    try:
        import matplotlib  # noqa: F401
        import seaborn  # noqa: F401
        return True
    except ImportError:
        return False


def _setup():
    import matplotlib
    matplotlib.use('Agg')                       # без окон: рисуем в файлы
    import matplotlib.pyplot as plt
    import seaborn as sns
    sns.set_theme(style='whitegrid', rc={
        'figure.facecolor': SURFACE, 'axes.facecolor': SURFACE, 'savefig.facecolor': SURFACE,
        'axes.edgecolor': BASELINE, 'axes.labelcolor': INK_2, 'axes.titlecolor': INK,
        'axes.titlesize': 13, 'axes.titleweight': 'semibold', 'axes.titlelocation': 'left',
        'axes.titlepad': 12, 'grid.color': GRID, 'grid.linewidth': 0.8, 'grid.linestyle': '-',
        'xtick.color': MUTED, 'ytick.color': INK_2, 'xtick.labelsize': 9.5, 'ytick.labelsize': 10.5,
        'font.family': ['Segoe UI', 'DejaVu Sans', 'Arial', 'sans-serif'], 'legend.frameon': False,
    })
    return plt, sns


def _num(v):
    """1 284 / 12,9 тыс. / 4,2 млн — в подписях."""
    if v is None:
        return ''
    a = abs(v)
    if a >= 1_000_000:
        return f'{v / 1_000_000:.1f} млн'.replace('.', ',')
    if a >= 10_000:
        return f'{v / 1_000:.1f}'.replace('.', ',') + ' тыс.'
    if float(v).is_integer() or a >= 100:
        return f'{v:,.0f}'.replace(',', ' ')
    return f'{v:.1f}'.replace('.', ',')


def _tick(v, _=None):
    """Подписи оси — круглые: 10 тыс. вместо 10,0 тыс., чтобы не слипались."""
    if abs(v) >= 1_000_000:
        return f'{v / 1_000_000:g} млн'.replace('.', ',')
    if abs(v) >= 10_000:
        return f'{v / 1_000:g}'.replace('.', ',') + ' тыс.'
    return f'{v:,.0f}'.replace(',', ' ')


def _short(text, limit=38, site=''):
    """Подпись полосы. Адрес страницы — без протокола и домена сайта:
    «msk · /blog/kak-vybrat-…» вместо полного URL."""
    t = str(text)
    if t.startswith(('http://', 'https://')):
        from urllib.parse import urlparse
        u = urlparse(t)
        host = u.netloc[4:] if u.netloc.startswith('www.') else u.netloc
        path = u.path or '/'
        site = (site or '').lower().replace('www.', '')
        if site and host == site:
            t = path
        elif site and host.endswith('.' + site):
            t = f'{host[:-len(site) - 1]} · {path}'
        else:
            t = host + path
    return t if len(t) <= limit else t[:limit - 1] + '…'


def _figure(plt, width, height, title, note=None):
    """Холст фиксированной ширины: подписи ужимают область графика, а не
    расширяют картинку. Тогда все графики панели в одном масштабе — на узкой
    карточке шрифт не мельчает из-за длинных адресов страниц."""
    fig, ax = plt.subplots(figsize=(width, height), layout='constrained')
    fig.suptitle(title, x=0.01, ha='left', fontsize=13, fontweight='semibold', color=INK)
    if note:
        import textwrap
        fig.supxlabel(textwrap.fill(note, int(width * 15)), x=0.01, ha='left', fontsize=8.5, color=MUTED)
    return fig, ax


def _save(fig, out_dir, name):
    path = out_dir / f'{name}.png'
    fig.savefig(path, dpi=130)
    import matplotlib.pyplot as plt
    plt.close(fig)
    return path


def _bars(plt, sns, labels, values, title, out_dir, name, fmt=_num, color=ACCENT, note=None):
    """Горизонтальные полосы одного цвета, значение — на конце полосы."""
    # Полосы — напрямую через matplotlib: seaborn усреднил бы строки с одинаковой
    # подписью (две длинные страницы после сокращения), и график бы соврал
    n = len(labels)
    fig, ax = _figure(plt, 6.4, 0.44 * n + 1.2 + (0.35 if note else 0), title, note)
    ax.barh(range(n), values, color=color, height=0.6)
    ax.set_yticks(range(n))
    ax.set_yticklabels(labels)
    ax.invert_yaxis()
    ax.set_xlabel('')
    ax.set_ylabel('')
    ax.grid(axis='y', visible=False)
    ax.xaxis.grid(True)
    sns.despine(ax=ax, left=False, bottom=True)
    ax.tick_params(axis='x', length=0)
    ax.tick_params(axis='y', length=0)
    top = max(values) if values else 0
    ax.set_xlim(0, top * 1.24 if top else 1)    # запас справа под подпись «32,3 тыс.»
    for i, v in enumerate(values):
        ax.text(v + top * 0.012, i, fmt(v), va='center', ha='left', fontsize=10, color=INK_2)
    ax.xaxis.set_major_locator(plt.MaxNLocator(5))
    ax.xaxis.set_major_formatter(plt.FuncFormatter(_tick))
    return _save(fig, out_dir, name)


def _dynamics(plt, sns, dates, cur, prev, group_label, out_dir, partial_tail=0, title=None, name='dynamics',
              fmt=None, prev_label='прошлый период', cur_label='текущий период', goal_cur=None, goal_prev=None,
              goal_label='Достижения целей'):
    """Визиты текущего периода и прошлого — на одной оси, прошлый серым.

    partial_tail — сколько последних точек Метрика, похоже, ещё не досчитала:
    они идут пунктиром с подписью, чтобы недосчёт не читался как обвал трафика.

    goal_cur/goal_prev — достижения главных целей: отдельной полосой под визитами с общей
    осью дат. На одну ось с визитами их не положить — визитов в сотни раз больше, а
    вторая ось врёт масштабом."""
    if goal_cur and any(goal_cur):
        fig, (ax, ag) = plt.subplots(2, 1, figsize=(11, 5.9), sharex=True, layout='constrained',
                                     gridspec_kw={'height_ratios': [2.2, 1]})
        fig.suptitle(title or f'Визиты по {group_label}', x=0.01, ha='left', fontsize=13, fontweight='semibold',
                     color=INK)
    else:
        goal_cur = None
        fig, ax = _figure(plt, 11, 3.8, title or f'Визиты по {group_label}')
    fmt = fmt or _num
    x = list(range(len(cur)))
    if prev and any(v is not None for v in prev):
        px = x[:len(prev)]
        ax.plot(px, prev[:len(px)], color=PREVIOUS, linewidth=2, solid_capstyle='round', label=prev_label)
    k = max(1, len(cur) - partial_tail)          # до этой точки данные полные
    ax.plot(x[:k], cur[:k], color=ACCENT, linewidth=2, solid_capstyle='round', label=cur_label)
    ax.fill_between(x[:k], cur[:k], color=ACCENT, alpha=0.08, linewidth=0)
    if partial_tail:
        ax.plot(x[k - 1:], cur[k - 1:], color=ACCENT, linewidth=2, linestyle=(0, (1.5, 2.5)), alpha=0.6)
        ax.scatter([x[-1]], [cur[-1]], s=46, color=SURFACE, edgecolor=ACCENT, linewidth=2, zorder=3)
        ax.annotate(f'{fmt(cur[-1])} — данные, вероятно, неполные', (x[-1], cur[-1]), xytext=(-10, 4),
                    textcoords='offset points', va='bottom', ha='right', fontsize=9.5, color=INK_2)
    else:
        ax.scatter([x[-1]], [cur[-1]], s=46, color=ACCENT, edgecolor=SURFACE, linewidth=2, zorder=3)
        ax.annotate(fmt(cur[-1]), (x[-1], cur[-1]), xytext=(8, 0), textcoords='offset points',
                    va='center', fontsize=9.5, color=INK_2)
    step = max(1, len(x) // 10)
    ax.set_xticks(x[::step])
    ax.set_xticklabels([d.strftime('%d.%m') for d in dates[::step]])
    ax.set_ylim(bottom=0)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(_tick))
    ax.grid(axis='x', visible=False)
    sns.despine(ax=ax, left=True)
    ax.tick_params(length=0)
    # легенда — над графиком справа: внутри она ложится на линию, когда значения у верхнего края
    ax.legend(loc='lower right', bbox_to_anchor=(1, 1.01), borderaxespad=0, fontsize=9.5, ncol=2, labelcolor=INK_2)
    if goal_cur:
        gx = list(range(len(goal_cur)))
        if goal_prev and any(goal_prev):
            ag.plot(gx[:len(goal_prev)], goal_prev[:len(gx)], color=PREVIOUS, linewidth=1.8, linestyle=(0, (4, 3)),
                    label=prev_label)
        gk = max(1, len(goal_cur) - partial_tail)
        ag.plot(gx[:gk], goal_cur[:gk], color=GOAL_C, linewidth=2, solid_capstyle='round', label=cur_label)
        if partial_tail:
            ag.plot(gx[gk - 1:], goal_cur[gk - 1:], color=GOAL_C, linewidth=2, linestyle=(0, (1.5, 2.5)), alpha=0.6)
        ag.scatter([gx[-1]], [goal_cur[-1]], s=36, color=GOAL_C, edgecolor=SURFACE, linewidth=1.5, zorder=3)
        ag.annotate(_num(goal_cur[-1]), (gx[-1], goal_cur[-1]), xytext=(8, 0), textcoords='offset points',
                    va='center', fontsize=9.5, color=INK_2)
        ag.set_title(goal_label, loc='left', fontsize=10.5, color=INK_2, pad=4, fontweight='normal')
        ag.set_ylim(bottom=0)
        ag.yaxis.set_major_formatter(plt.FuncFormatter(_tick))
        ag.grid(axis='x', visible=False)
        sns.despine(ax=ag, left=True)
        ag.tick_params(length=0)
        ag.legend(loc='lower right', bbox_to_anchor=(1, 1.01), borderaxespad=0, fontsize=9, ncol=2,
                  labelcolor=INK_2)
    return _save(fig, out_dir, name)


def _pair(plt, sns, labels, visits, reach, conv, title, out_dir, name, goal_label='Достижения целей', note=None):
    """Рейтинг в две колонки с общими подписями: слева визиты, справа достижения главных
    целей и конверсия в них. У каждой колонки свой масштаб — иначе достижения не видны."""
    n = len(labels)
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(6.4, 0.44 * n + 1.5 + (0.35 if note else 0)), sharey=True,
                                 layout='constrained', gridspec_kw={'width_ratios': [1.15, 1]})
    fig.suptitle(title, x=0.01, ha='left', fontsize=13, fontweight='semibold', color=INK)
    if note:
        import textwrap
        fig.supxlabel(textwrap.fill(note, 96), x=0.01, ha='left', fontsize=8.5, color=MUTED)
    y = list(range(n))
    a1.barh(y, [v or 0 for v in visits], color=ACCENT, height=0.6)
    a2.barh(y, [r or 0 for r in reach], color=GOAL_C, height=0.6)
    a1.set_yticks(y)
    a1.set_yticklabels(labels)
    a1.invert_yaxis()
    pct = lambda c: f'{c:.1f}%'.replace('.', ',')  # noqa: E731
    texts = ([_num(v) for v in visits],
             [f'{_num(r)} · {pct(c)}' if c is not None else _num(r) for r, c in zip(reach, conv)])
    for ax, vals, txt, head in ((a1, visits, texts[0], 'Визиты'), (a2, reach, texts[1], goal_label)):
        top = max((v or 0 for v in vals), default=0)
        ax.set_xlim(0, top * 1.55 if top else 1)
        for i, (v, t) in enumerate(zip(vals, txt)):
            ax.text((v or 0) + top * 0.02, i, t, va='center', ha='left', fontsize=9.5, color=INK_2)
        ax.set_title(head, loc='left', fontsize=10.5, color=INK_2, pad=4, fontweight='normal')
        ax.grid(False)
        ax.xaxis.set_visible(False)
        sns.despine(ax=ax, left=ax is a2, bottom=True)
        ax.tick_params(length=0)
    return _save(fig, out_dir, name)


def _changes(plt, sns, kpis, out_dir, title='Что изменилось к прошлому периоду', name='changes',
             default_good=None, note='Зелёным — стало лучше, красным — хуже. '
                                     'Для долей в процентах — разница в процентных пунктах.'):
    """Что изменилось к прошлому периоду: зелёным — лучше, красным — хуже.

    default_good — для подписей не из UP_IS_GOOD (источники трафика: рост — хорошо)."""
    rows = [(label, ch, unit) for label, _, _, ch, unit in kpis if ch is not None]
    if not rows:
        return None
    fig, ax = _figure(plt, 6.4, 0.44 * len(rows) + 1.55, title, note)
    labels = [lbl(r[0]) for r in rows]
    vals = [r[1] for r in rows]
    colors = []
    for label, ch, _ in rows:
        good = UP_IS_GOOD.get(label, default_good)
        colors.append(MUTED if good is None or ch == 0 else (GOOD if (ch > 0) == good else BAD))
    ax.barh(range(len(rows)), vals, color=colors, height=0.6)
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels(labels)
    ax.invert_yaxis()
    ax.axvline(0, color=BASELINE, linewidth=1)
    lim = max(abs(v) for v in vals) or 1
    ax.set_xlim(-lim * 1.45, lim * 1.45)
    for i, (label, ch, unit) in enumerate(rows):
        txt = f'{ch:+.1f}'.replace('.', ',') + (' п.п.' if unit == 'pp' else '%')
        ax.text(ch + (lim * 0.03 if ch >= 0 else -lim * 0.03), i, txt, va='center',
                ha='left' if ch >= 0 else 'right', fontsize=10, color=INK_2)
    ax.grid(axis='y', visible=False)
    ax.xaxis.set_visible(False)
    sns.despine(ax=ax, left=True, bottom=True)
    ax.tick_params(length=0)
    return _save(fig, out_dir, name)


# Источники в подписях тепловой карты — коротко, как в обиходе
SOURCE_SHORT = {
    'Переходы из поисковых систем': 'Поиск', 'Переходы по рекламе': 'Реклама', 'Прямые заходы': 'Прямые',
    'Внутренние переходы': 'Внутренние', 'Переходы по ссылкам на сайтах': 'Ссылки',
    'Переходы из социальных сетей': 'Соцсети', 'Переходы из рекомендательных систем': 'Рекомендации',
    'Переходы с сохранённых страниц': 'Сохранённые', 'Переходы из мессенджеров': 'Мессенджеры',
    'Переходы из почтовых систем': 'Почта',
}
WIDE = {'dynamics', 'goal_sources', 'cost_dynamics', 'goal_campaigns', 'weeks', 'goal_engines', 'hours'}   # графики на всю ширину панели


def _pct(v):
    if not v:
        return '0'
    return '<0,1' if v < 0.1 else f'{v:.1f}'.replace('.', ',')


def _goal_sources(plt, sns, gs, out_dir, by='источники', min_visits=100, name='goal_sources',
                  title='Конверсия в цели по источникам, %'):
    """Конверсия в каждую цель по источникам: строки — цели, столбцы — источники.

    Цели разного масштаба: «3 страницы» — десятки процентов, звонок — доли
    процента. Общая шкала залила бы мелкие цели одним бледным цветом, поэтому
    цвет сравнивает источники внутри строки: темнее — где у этой цели конверсия
    выше. Число в клетке — сама конверсия. by='кампании' — та же карта по
    кампаниям Директа."""
    import textwrap
    import numpy as np
    from matplotlib.colors import LinearSegmentedColormap
    goals = [g for g in gs['goals'] if g[2]][:10]
    sources = [s for s in gs['sources'] if (s[1] or 0) >= min_visits][:7]
    if not goals or len(sources) < 2:
        return None
    conv = np.array([[((gs['cells'].get((gid, src)) or (None, None))[1] or 0) for src, _ in sources]
                     for gid, _, _ in goals], dtype=float)
    top = conv.max(axis=1, keepdims=True)
    top[top == 0] = 1
    labels = np.array([[_pct(v) for v in row] for row in conv])
    note = (f'Цвет сравнивает {by} внутри каждой цели: темнее — там, где у этой цели конверсия выше. '
            f'{by.capitalize()} от {min_visits} визитов; у цели — число достижений за период.')
    fig, ax = _figure(plt, 11, 0.5 * len(goals) + 2.3, title, note)
    sns.heatmap(conv / top, annot=labels, fmt='', cbar=False, vmin=0, vmax=1, ax=ax,
                cmap=LinearSegmentedColormap.from_list('accent', ['#eef4fc', ACCENT]),
                linewidths=2, linecolor=SURFACE, annot_kws={'fontsize': 10})
    ax.set_xticklabels([textwrap.fill(SOURCE_SHORT.get(src) or _short(src or '(не определено)', 32), 17)
                        + f'\n{_num(v)} виз.' for src, v in sources],
                       rotation=0, fontsize=9.5)
    ax.set_yticklabels([f'{_short(name, 34)} · {_num(t)}' for _, name, t in goals], rotation=0)
    ax.xaxis.tick_top()
    ax.tick_params(length=0)
    ax.set_xlabel('')
    ax.set_ylabel('')
    return _save(fig, out_dir, name)


def build(ctx, out_dir):
    """Рисует графики по собранным данным. Возвращает [(заголовок, путь, подпись)]."""
    plt, sns = _setup()
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = {}
    t = ctx['tables']

    dyn = ctx.get('dynamics')
    if dyn and dyn['cur']:
        p = _dynamics(plt, sns, dyn['dates'], dyn['cur'], dyn.get('prev'), dyn['label'], out_dir,
                      dyn.get('partial_tail', 0), prev_label=ctx.get('prev_label', 'прошлый период'),
                      cur_label=ctx.get('cur_label', 'текущий период'), goal_cur=dyn.get('goal_cur'),
                      goal_prev=dyn.get('goal_prev'), goal_label=lbl('Достижения целей'))
        charts['dynamics'] = ('Динамика визитов', p, f'Серым — {ctx.get("prev_label", "прошлый период")}, '
                                                     f'по порядку дней — лист «Динамика»')

    p = _changes(plt, sns, ctx['kpis'], out_dir, title=f'Что изменилось {ctx.get("vs", "к прошлому периоду")}')
    if p:
        charts['changes'] = ('Что изменилось', p, 'Лист «Сводка»')

    def rows_of(key, n, skip_undefined=False, min_visits=0, metric=0):
        if key not in t:
            return [], []
        dims, metrics, rows, _ = t[key]
        out = [(r[0][0], r[1][metric]) for r in rows
               if r[1][metric] is not None and (not skip_undefined or r[0][0] not in (None, ''))
               and (r[1][0] or 0) >= min_visits]
        out = out[:n]
        return ([_short(x[0] if x[0] not in (None, '') else '(не определено)', site=ctx['site']) for x in out],
                [x[1] for x in out])

    def goal_idx(key):
        ids = [m[0] for m in t[key][1]]
        return (ids.index('ym:s:sumGoalReachesAny') if 'ym:s:sumGoalReachesAny' in ids else None,
                ids.index('ym:s:anyGoalConversionRate') if 'ym:s:anyGoalConversionRate' in ids else None)

    for key, title, name, caption, kw in (
        ('sources', 'Визиты по источникам', 'sources', 'Последний значимый источник — лист «Источники»', {}),
        ('engines', 'Визиты из поисковых систем', 'engines', 'Лист «Поисковые системы»', {}),
        ('entry', 'Страницы входа — топ-10', 'entry', 'С каких страниц начинались визиты — лист «Страницы входа»', {}),
        ('phrases', 'Поисковые фразы — топ-12', 'phrases',
         'Большую часть фраз поисковики скрывают — лист «Поисковые фразы»', {'skip_undefined': True}),
        ('devices', 'Визиты по устройствам', 'devices', 'Лист «Устройства»', {}),
        ('cities', 'Города — топ-10', 'cities', 'Лист «Города»', {'skip_undefined': True}),
    ):
        n = 12 if key == 'phrases' else 10
        labels, values = rows_of(key, n, **kw)
        if len(labels) < 2:
            continue
        ri, ci = goal_idx(key) if ctx.get('focus') else (None, None)
        if ri is None:
            charts[name] = (title, _bars(plt, sns, labels, values, title, out_dir, name), caption)
            continue
        # те же строки, что у визитов, — достижения и конверсия главных целей рядом
        sel = [r for r in t[key][2] if r[1][0] is not None
               and (not kw.get('skip_undefined') or r[0][0] not in (None, ''))][:n]
        charts[name] = (title, _pair(plt, sns, labels, values, [r[1][ri] for r in sel],
                                     [r[1][ci] if ci is not None else None for r in sel], title, out_dir, name,
                                     goal_label=lbl('Достижения целей')), caption)

    # Конверсия по источникам — только источники от 100 визитов, иначе шум
    if 'sources' in t:
        dims, metrics, rows, _ = t['sources']
        idx = next((i for i, m in enumerate(metrics) if m[0] == 'ym:s:anyGoalConversionRate'), None)
        if idx is not None:
            sel = [(r[0][0], r[1][idx]) for r in rows if (r[1][0] or 0) >= 100 and r[1][idx] is not None]
            sel.sort(key=lambda x: -x[1])
            if len(sel) >= 2:
                p = _bars(plt, sns, [_short(s[0], site=ctx['site']) for s in sel], [s[1] for s in sel],
                          lbl('Конверсия в любую цель, %')[:-3] + ' по источникам, %', out_dir, 'conversion',
                          fmt=lambda v: f'{v:.1f}%'.replace('.', ','),
                          note='Источники от 100 визитов.' + ('' if ctx.get('focus') else
                               ' Цели вовлечённости («3 страницы», «время на сайте») тоже входят в конверсию.'))
                charts['conversion'] = ('Конверсия по источникам', p, 'Лист «Источники»')

    goals = ctx.get('goals') or []
    goals = [g for g in goals if g[1]][:10]
    if len(goals) >= 2:
        p = _bars(plt, sns, [_short(g[0]) for g in goals], [g[1] for g in goals],
                  'Цели — топ-10 по достижениям', out_dir, 'goals')
        charts['goals'] = ('Цели', p, 'Достижения и конверсия каждой цели — лист «Цели»')
    # Порядок на панели: соседние графики в паре дополняют друг друга
    gs = ctx.get('goal_sources')
    if gs:
        p = _goal_sources(plt, sns, gs, out_dir)
        if p:
            charts['goal_sources'] = ('Цели по источникам', p,
                                      'Достижения и доля каждого источника в цели — лист «Цели по источникам»')
    order = ['dynamics', 'changes', 'sources', 'conversion', 'engines', 'devices',
             'entry', 'phrases', 'cities', 'goals', 'goal_sources']
    return [charts[k] for k in order if k in charts]


# Для каких показателей Директа рост — это хорошо
UP_IS_GOOD.update({'Клики': True, 'Расход, ₽': None, 'Цена клика, ₽': False, 'Цена цели, ₽': False})


def _money(v):
    return f'{_num(v)} ₽' if v is not None else ''


def _split(plt, sns, labels, visits, conv, title, out_dir, name, width=6.4, note=None, cost=None, cpa=None):
    """Графики друг под другом: визиты, расход и цена цели (если есть расходы) и
    отклонение конверсии от средней.

    Конверсия везде близка (цели вовлечённости поднимают её до 80–90 %), и столбики
    от нуля разницы не покажут; отклонение в п.п. показывает, когда реклама работает
    лучше среднего (зелёным) и хуже (красным). Средняя — с весом по визитам."""
    total = sum(v or 0 for v in visits)
    avg = sum((c or 0) * (v or 0) for c, v in zip(conv, visits)) / total if total else 0
    dev = [(c - avg) if c is not None and v else 0 for c, v in zip(conv, visits)]
    money = bool(cost) and any(cost)
    fig, axes = plt.subplots(4 if money else 2, 1, figsize=(width, (8.8 if money else 4.8) + (0.3 if note else 0)),
                             sharex=True, layout='constrained')
    a1, a2 = axes[0], axes[-1]
    fig.suptitle(title, x=0.01, ha='left', fontsize=13, fontweight='semibold', color=INK)
    if note:
        import textwrap
        fig.supxlabel(textwrap.fill(note, int(width * 15)), x=0.01, ha='left', fontsize=8.5, color=MUTED)
    x = list(range(len(labels)))
    a1.bar(x, [v or 0 for v in visits], color=ACCENT, width=0.7)
    a1.set_title('Визиты', loc='left', fontsize=10.5, color=INK_2, pad=4, fontweight='normal')
    a1.yaxis.set_major_formatter(plt.FuncFormatter(_tick))
    if money:
        ac, aq = axes[1], axes[2]
        ac.bar(x, [c or 0 for c in cost], color=ACCENT, width=0.7)
        ac.set_title('Расход, ₽', loc='left', fontsize=10.5, color=INK_2, pad=4, fontweight='normal')
        ac.yaxis.set_major_formatter(plt.FuncFormatter(_tick))
        # средняя цена цели — весь расход ÷ все достижения; дешевле неё — зелёным
        reach = sum((c or 0) / q for c, q in zip(cost, cpa or []) if c and q)
        avg_cpa = sum(c or 0 for c, q in zip(cost, cpa or []) if c and q) / reach if reach else 0
        qs = [q or 0 for q in (cpa or [None] * len(x))]
        aq.bar(x, qs, color=[MUTED if not q else GOOD if q <= avg_cpa else BAD for q in qs], width=0.7)
        if avg_cpa:
            aq.axhline(avg_cpa, color=INK_2, linewidth=1, linestyle=(0, (3, 3)))
        aq.set_title(f'Цена цели, ₽: дешевле средней {_num(round(avg_cpa, 2))} ₽ — зелёным', loc='left',
                     fontsize=10.5, color=INK_2, pad=4, fontweight='normal')
        aq.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: _num(v)))
    a2.bar(x, dev, color=[GOOD if d > 0 else BAD for d in dev], width=0.7)
    a2.axhline(0, color=BASELINE, linewidth=1)
    a2.set_title(f'Конверсия: отклонение от средней {_num(round(avg, 1))}%, п.п.', loc='left', fontsize=10.5,
                 color=INK_2, pad=4, fontweight='normal')
    lim = max((abs(d) for d in dev), default=1) or 1
    a2.set_ylim(-lim * 1.25, lim * 1.25)
    for ax in axes:
        ax.grid(axis='x', visible=False)
        sns.despine(ax=ax, left=True, bottom=True)
        ax.tick_params(length=0)
    a2.set_xticks(x)
    a2.set_xticklabels(labels, fontsize=9.5 if len(labels) > 12 else 10)
    return _save(fig, out_dir, name)


def build_direct(ctx, out_dir):
    """Графики отчёта по Директу. Возвращает [(заголовок, путь, подпись)]."""
    plt, sns = _setup()
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = {}
    t = ctx['tables']

    dyn = ctx.get('dynamics')
    if dyn and dyn['cur']:
        p = _dynamics(plt, sns, dyn['dates'], dyn['cur'], dyn.get('prev'), dyn['label'], out_dir,
                      dyn.get('partial_tail', 0), title=f'Визиты из Директа по {dyn["label"]}',
                      prev_label=ctx.get('prev_label', 'прошлый период'),
                      cur_label=ctx.get('cur_label', 'текущий период'), goal_cur=dyn.get('goal_cur'),
                      goal_prev=dyn.get('goal_prev'), goal_label=lbl('Достижения целей'))
        charts['dynamics'] = ('Визиты из Директа', p, 'Прошлый период совмещён по порядку дней — лист «Динамика»')
        if dyn.get('cost') and any(dyn['cost']):
            p = _dynamics(plt, sns, dyn['dates'], dyn['cost'], None, dyn['label'], out_dir,
                          dyn.get('partial_tail', 0), title=f'Расход по {dyn["label"]}, ₽', name='cost_dynamics',
                          fmt=_money, cur_label='расход')
            charts['cost_dynamics'] = ('Расход', p, 'Расход — по данным Директа, лист «Динамика»')

    p = _changes(plt, sns, ctx['kpis'], out_dir, title=f'Что изменилось {ctx.get("vs", "к прошлому периоду")}')
    if p:
        charts['changes'] = ('Что изменилось', p, 'Лист «Сводка»')

    def column(key, label):
        if key not in t:
            return None, None, None
        heads, mets, rows, _ = t[key]
        names = [m[1] for m in mets]
        return (names.index(label) if label in names else None), rows, len(heads)

    def bars(key, label, title, name, caption, n=10, min_visits=0, fmt=_num, note=None, reverse=False):
        i, rows, nd = column(key, label)
        if i is None:
            return
        sel = [(r[0][0], r[1][i]) for r in rows if r[1][i] is not None and (r[1][0] or 0) >= min_visits]
        if label != 'Визиты':
            sel.sort(key=lambda x: x[1], reverse=not reverse)
        sel = sel[:n]
        if len(sel) >= 2:
            labels = [_short(x[0] if x[0] not in (None, '') else '(не определено)', site=ctx['site']) for x in sel]
            names = [m[1] for m in t[key][1]]
            if label == 'Визиты' and ctx.get('focus') and 'Достижения целей' in names:
                # визиты и рядом — достижения главных целей с конверсией, по тем же строкам
                ri = names.index('Достижения целей')
                ci = names.index('Конверсия в любую цель, %') if 'Конверсия в любую цель, %' in names else None
                rows_sel = [r for r in rows if r[1][i] is not None and (r[1][0] or 0) >= min_visits][:n]
                charts[name] = (title, _pair(plt, sns, labels, [x[1] for x in sel],
                                             [r[1][ri] for r in rows_sel],
                                             [r[1][ci] if ci is not None else None for r in rows_sel],
                                             title, out_dir, name, goal_label=lbl('Достижения целей'),
                                             note=note), caption)
                return
            charts[name] = (title, _bars(plt, sns, labels, [x[1] for x in sel], title, out_dir, name,
                                         fmt=fmt, note=note), caption)

    pct = lambda v: f'{v:.1f}%'.replace('.', ',')  # noqa: E731
    # Кампании — из Директа, а если Яндекс их не отдал — по метке utm_campaign
    ck, csheet = ('campaigns', 'Кампании') if 'campaigns' in t else ('utm_campaign', 'Кампании по UTM')
    ctitle = 'кампаниям' if ck == 'campaigns' else 'кампаниям (utm_campaign)'
    bars(ck, 'Визиты', f'Визиты по {ctitle}', 'campaigns', f'Лист «{csheet}»')
    bars(ck, 'Конверсия в любую цель, %', f'Конверсия по {ctitle}, %', 'camp_conv', f'Лист «{csheet}»',
         min_visits=30, fmt=pct, note='Кампании от 30 визитов.' + ('' if ctx.get('focus') else
                                      ' Цели вовлечённости тоже входят в конверсию.'))

    # Автотаргетинг против ключевых фраз — по метке utm_term
    if t.get('autotarget'):
        _, mets, rows, _ = t['autotarget']
        names = [m[1] for m in mets]
        ci = names.index('Конверсия в любую цель, %')
        # без метки utm_term — обычно в разы больше и забивает масштаб: только в подписи
        untagged = sum(r[1][0] or 0 for r in rows if r[0][0] == 'Без метки utm_term')
        sel = [(r[0][0], r[1][0], r[1][ci]) for r in rows if r[1][0] and r[0][0] != 'Без метки utm_term']
        if sel:
            labels = [f'{n} · {pct(c)}' if c is not None else n for n, _, c in sel]
            charts['autotarget'] = ('Автотаргетинг и фразы', _bars(
                plt, sns, labels, [v for _, v, _ in sel], 'Автотаргетинг и ключевые фразы: визиты', out_dir,
                'autotarget', note='После точки — конверсия. По метке utm_term: «---autotargeting» '
                                   '— автотаргетинг, остальное — ключевые фразы.'
                                   + (f' Ещё {_num(untagged)} визитов — без метки utm_term (кампании, где фраза '
                                      f'в метку не подставляется).' if untagged else '')),
                'Лист «Автотаргетинг и фразы»; сами фразы — «Фразы по UTM»')

    # Время и аудитория: визиты и отклонение конверсии от средней
    def split(key, title, name, short=None, width=6.4, extra=()):
        parts = [k for k in (key, *extra) if t.get(k)]
        labels, visits, conv, cost, cpa = [], [], [], [], []
        for k in parts:
            _, mets, rows, _ = t[k]
            names = [m[1] for m in mets]
            if 'Конверсия в любую цель, %' not in names:
                continue                     # лист только из расходов — визитов нет
            ci = names.index('Конверсия в любую цель, %')
            ri = names.index('Расход, ₽') if 'Расход, ₽' in names else None
            qi = names.index('Цена цели, ₽') if 'Цена цели, ₽' in names else None
            for r in rows:
                labels.append(short(r[0][0]) if short else str(r[0][0]))
                visits.append(r[1][0])
                conv.append(r[1][ci])
                cost.append(r[1][ri] if ri is not None else None)
                cpa.append(r[1][qi] if qi is not None else None)
        if len(labels) >= 2:
            charts[name] = (title, _split(plt, sns, labels, visits, conv, title, out_dir, name, width=width,
                                          cost=cost, cpa=cpa),
                            ('Листы «' if len(parts) > 1 else 'Лист «') + '», «'.join(SHEETS.get(k, k) for k in parts) + '»')

    SHEETS = {'hours': 'Часы', 'weekdays': 'Дни недели', 'age': 'Возраст', 'gender': 'Пол',
              'newold': 'Новые и вернувшиеся'}
    days = {'Понедельник': 'Пн', 'Вторник': 'Вт', 'Среда': 'Ср', 'Четверг': 'Чт', 'Пятница': 'Пт',
            'Суббота': 'Сб', 'Воскресенье': 'Вс'}
    who = {'мужской': 'Мужчины', 'женский': 'Женщины', 'Новые посетители': 'Новые',
           'Вернувшиеся посетители': 'Вернувшиеся'}
    split('hours', 'Реклама по часам суток', 'hours', short=lambda h: str(h)[:2], width=11)
    split('weekdays', 'Реклама по дням недели', 'weekdays', short=lambda d: days.get(d, d))
    split('age', 'Возраст', 'age', short=lambda a: str(a).replace(' года', '').replace(' лет и старше', '+')
          .replace('младше 18 лет', '<18'))
    split('gender', 'Пол и новые посетители', 'audience', short=lambda g: who.get(g, g), extra=('newold',))
    if ctx.get('costs'):
        bars('campaigns', 'Расход, ₽', 'Расход по кампаниям, ₽', 'camp_cost', 'Лист «Кампании»', fmt=_money)
        bars('campaigns', 'Цена цели, ₽', 'Цена цели по кампаниям, ₽', 'camp_cpa',
             'Расход кампании, делённый на достижения главных целей — лист «Кампании»', fmt=_money, reverse=True,
             note='Дешевле — выше.' + ('' if ctx.get('focus') else ' Цели вовлечённости тоже входят в достижения.'))
    bars('placetype', 'Визиты', 'Поиск и сети: визиты', 'placetype', 'Лист «Поиск и сети»')
    bars('placetype', 'Расход, ₽', 'Поиск и сети: расход, ₽', 'placetype_cost', 'Лист «Поиск и сети»', fmt=_money)
    if t.get('cities'):
        _, mets, rows, _ = t['cities']
        names = [m[1] for m in mets]
        if 'Расход, ₽' in names:
            ri, qi = names.index('Расход, ₽'), names.index('Цена цели, ₽')
            top = sorted((r for r in rows if r[1][ri]), key=lambda r: -r[1][ri])[:10]
            if len(top) >= 2:
                charts['cities_cost'] = ('Города в деньгах', _bars(
                    plt, sns, [f'{r[0][0]} · цель {_money(r[1][qi])}' if r[1][qi] else str(r[0][0]) for r in top],
                    [r[1][ri] for r in top], 'Города: расход, ₽ — топ-10', out_dir, 'cities_cost', fmt=_money,
                    note='После точки — цена цели в городе.'), 'Лист «Города»')
    if t.get('goal_cost'):
        top = sorted(t['goal_cost'], key=lambda g: -g[1])[:12]
        if top:
            charts['goal_cost'] = ('Цена целей', _bars(
                plt, sns, [f'{_short(n, 34)} · {_num(r)}' for n, r, _, _ in top], [c for _, _, c, _ in top],
                'Цена цели, ₽ — по каждой цели', out_dir, 'goal_cost', fmt=_money,
                note='Весь расход Директа за период, делённый на достижения цели (после точки).'), 'Лист «Цена целей»')
    bars('conditions', 'Визиты', 'Ключевые фразы — топ-12 по визитам', 'conditions', 'Лист «Ключевые фразы»', n=12)
    bars('queries', 'Визиты', 'Поисковые запросы — топ-12', 'queries', 'Что искали перед кликом — лист «Поисковые запросы»',
         n=12)
    bars('entry', 'Визиты', 'Страницы входа из рекламы — топ-10', 'entry', 'Лист «Страницы входа»')

    native = bool((ctx.get('goal_sources') or {}).get('goals'))
    gs = ctx.get('goal_sources') if native else ctx.get('goal_sources_utm')
    if gs:
        p = _goal_sources(plt, sns, gs, out_dir, by='кампании', min_visits=30, name='goal_campaigns',
                          title='Конверсия в цели по кампаниям, %')
        if p:
            charts['goal_campaigns'] = ('Цели по кампаниям', p,
                                        'Достижения, доля и цена цели по каждой кампании — лист «Цели по кампаниям'
                                        + ('»' if native else ' UTM»'))
    order = ['dynamics', 'cost_dynamics', 'changes', 'goal_cost', 'campaigns', 'camp_conv', 'camp_cost', 'camp_cpa',
             'goal_campaigns', 'autotarget', 'placetype', 'placetype_cost', 'hours', 'weekdays', 'age', 'audience',
             'cities_cost', 'entry', 'conditions', 'queries']
    return [charts[k] for k in order if k in charts]


def build_seo(ctx, out_dir):
    """Графики SEO-отчёта: поиск к сравниваемому периоду, страницы, длинные клики, цели."""
    from datetime import date
    plt, sns = _setup()
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    charts = {}
    t = ctx['tables']
    vs, prev, cur = ctx.get('vs', 'к году назад'), ctx.get('prev_label', 'год назад'), ctx.get('cur_label', 'этот год')
    step = ctx.get('step', 'неделям')

    if t.get('weeks') and t['weeks'][0]:
        cw, pw = t['weeks']
        goals = len(cw[0]) > 4
        p = _dynamics(plt, sns, [date.fromisoformat(w[0]) for w in cw], [w[1] for w in cw],
                      [w[1] for w in pw] or None, step, out_dir,
                      title=f'Визиты из поиска по {step}: {cur} и {prev}', name='weeks',
                      prev_label=prev, cur_label=cur, goal_cur=[w[4] for w in cw] if goals else None,
                      goal_prev=[w[4] for w in pw] if goals and pw and len(pw[0]) > 4 else None,
                      goal_label=lbl('Достижения целей'))
        charts['weeks'] = ('Поиск в динамике', p, f'Серым — {prev}, по порядку {step} — лист «Поиск в динамике»')

    p = _changes(plt, sns, ctx['kpis'], out_dir, title=f'Поиск: что изменилось {vs}')
    if p:
        charts['changes'] = ('Что изменилось', p, 'Только визиты из поиска — лист «Сводка»')

    if t.get('sources'):
        rows = []
        for d, v, pv in t['sources']:
            if max(v[0] or 0, pv[0] or 0) < 30 or not pv[0] or v[0] is None:
                continue
            rows.append((SOURCE_SHORT.get(d[0], _short(d[0] or '(не определено)', 30)), None, None,
                         round((v[0] - pv[0]) / pv[0] * 100, 1), 'pct'))
        p = _changes(plt, sns, rows, out_dir, title=f'Визиты по источникам {vs}', name='sources_yoy',
                     default_good=True, note=f'Сравнение: {prev}. Источники от 30 визитов.')
        if p:
            charts['sources_yoy'] = ('Источники — сравнение', p, 'Лист «Источники — сравнение»')

    if t.get('entry'):
        mets, rows, was = t['entry']
        top = [r for r in rows if r[1][0]][:10]
        ids = [m[0] for m in mets]
        labels = [_short(r[0][0], site=ctx['site']) for r in top]
        if len(top) >= 2 and ctx.get('focus') and 'ym:s:sumGoalReachesAny' in ids:
            ri, ci = ids.index('ym:s:sumGoalReachesAny'), ids.index('ym:s:anyGoalConversionRate')
            charts['entry'] = ('Страницы входа из поиска', _pair(
                plt, sns, labels, [r[1][0] for r in top], [r[1][ri] for r in top], [r[1][ci] for r in top],
                'Страницы входа из поиска — топ-10', out_dir, 'entry', goal_label=lbl('Достижения целей')),
                'Лист «Страницы входа из поиска»')
        elif len(top) >= 2:
            charts['entry'] = ('Страницы входа из поиска', _bars(
                plt, sns, labels, [r[1][0] for r in top],
                'Страницы входа из поиска — топ-10', out_dir, 'entry'), 'Лист «Страницы входа из поиска»')
        now = {r[0][0]: r[1][0] or 0 for r in rows}
        lost = sorted(((u, pv - now.get(u, 0)) for u, pv in was.items() if pv and pv - now.get(u, 0) > 0),
                      key=lambda x: -x[1])[:10]
        if len(lost) >= 2:
            charts['entry_lost'] = ('Страницы, потерявшие поиск', _bars(
                plt, sns, [_short(x[0], site=ctx['site']) for x in lost], [x[1] for x in lost],
                f'Страницы, потерявшие визиты из поиска {vs}', out_dir, 'entry_lost', color=BAD,
                fmt=lambda v: f'−{_num(v)}', note=f'На сколько визитов из поиска меньше, чем в периоде «{prev}».'),
                f'Визиты сейчас и {prev} — лист «Страницы входа из поиска»')

    if t.get('long'):
        norm, least = ctx.get('long_norm', 10), ctx.get('long_min', 10)
        top = [r for r in t['long'] if r[1] >= least][:12]
        if len(top) >= 2:
            pct = lambda v: f'{v:.1f}%'.replace('.', ',')  # noqa: E731
            charts['long'] = ('Длинные клики', _bars(
                plt, sns, [_short(r[0], 40) for r in top], [r[3] or 0 for r in top],
                'Доля длинных кликов по фразам из Яндекса, %', out_dir, 'long', fmt=pct,
                color=[BAD if (r[3] or 0) < norm else ACCENT for r in top],
                note=f'Длинный визит — дольше {ctx.get("long_click", 59)} секунд. Норма — от {norm}–15 %: '
                     f'красным — ниже {norm} %. Фразы от {least} визитов, по убыванию визитов.'),
                'Все фразы и оценка — лист «Длинные клики»')

    gs = ctx.get('goal_sources')
    if gs:
        p = _goal_sources(plt, sns, gs, out_dir, by='поисковики', min_visits=30, name='goal_engines',
                          title='Конверсия в цели по поисковым системам, %')
        if p:
            charts['goal_engines'] = ('Цели из поиска', p, 'Достижения и доля каждой поисковой системы — лист «Цели из поиска»')
    order = ['weeks', 'changes', 'sources_yoy', 'entry', 'entry_lost', 'long', 'goal_engines']
    return [charts[k] for k in order if k in charts]


def _days(n):
    """14 дней, 1 день, 2 дня."""
    tail = 'дней' if 11 <= n % 100 <= 14 else {1: 'день', 2: 'дня', 3: 'дня', 4: 'дня'}.get(n % 10, 'дней')
    return f'{n} {tail}'


def _span(d1, d2):
    """28.09 – 11.10.2026; год пишется один раз, если он общий."""
    return (f'{d1:%d.%m} – {d2:%d.%m.%Y}' if d1.year == d2.year else f'{d1:%d.%m.%Y} – {d2:%d.%m.%Y}')


def period_line(ctx):
    """Строка периодов для Excel: «Период: 28.09 – 11.10.2026, 14 дней · сравнение: …»."""
    parts = [f'{name}: {_span(d1, d2)}, {_days((d2 - d1).days + 1)}' for name, d1, d2 in ctx['periods']]
    return ' · '.join(parts[:1] + [p[0].lower() + p[1:] for p in parts[1:]])


def find_browser():
    """Chrome или Edge — они печатают страницу отчёта в PDF. Нет ни того ни другого — None."""
    import os
    import shutil
    for name in ('chrome', 'google-chrome', 'google-chrome-stable', 'chromium', 'chromium-browser', 'msedge',
                 'microsoft-edge'):
        path = shutil.which(name)
        if path:
            return path
    local = os.environ.get('LOCALAPPDATA', '')
    for path in (r'C:\Program Files\Google\Chrome\Application\chrome.exe',
                 r'C:\Program Files (x86)\Google\Chrome\Application\chrome.exe',
                 local + r'\Google\Chrome\Application\chrome.exe',
                 r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
                 r'C:\Program Files\Microsoft\Edge\Application\msedge.exe',
                 '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome',
                 '/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge'):
        if path and Path(path).is_file():
            return path
    return None


def html_to_pdf(html_path):
    """PDF страницы отчёта: без шапки с логотипом и кнопок, со ссылками в колонтитуле
    (правила печати — PRINT_CSS). Печатает Chrome/Edge без окна, со своим временным профилем."""
    import shutil
    import subprocess
    import tempfile
    browser = find_browser()
    if not browser:
        return None
    html_path = Path(html_path).resolve()
    pdf = html_path.with_suffix('.pdf')
    if pdf.exists():
        pdf.unlink()                      # старый PDF не должен выдать себя за новый
    profile = tempfile.mkdtemp(prefix='metrika-pdf-')
    try:
        subprocess.run([browser, '--headless=new', '--disable-gpu', '--no-first-run', f'--user-data-dir={profile}',
                        '--no-pdf-header-footer', '--print-to-pdf-no-header', f'--print-to-pdf={pdf}',
                        '--virtual-time-budget=5000', html_path.as_uri()],
                       capture_output=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired):
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    return pdf if pdf.is_file() and pdf.stat().st_size > 0 else None


def print_css():
    """Печать в PDF: A4, без шапки с логотипом и кнопок, ссылки — в колонтитуле каждой страницы."""
    site = BRAND['site'].split('//')[-1].rstrip('/')
    tg = BRAND['telegram'][1].split('//')[-1].rstrip('/')
    foot = 'font:8.5pt system-ui,-apple-system,"Segoe UI",sans-serif;color:#6b6a65'
    return ('@page{size:A4;margin:12mm 11mm 16mm;'
            f'@bottom-left{{content:"{BRAND["logo_alt"]} · {site}";{foot}}}'
            f'@bottom-right{{content:"Telegram-канал: {tg} · стр. " counter(page) " из " counter(pages);{foot}}}}}'
            '@media print{.brand,.switch,.pdfbtn{display:none!important}body{background:#fff}'
            '.wrap{max-width:none;padding:0}.grid{grid-template-columns:1fr}'
            '.card,.tile,.periods,h2.sec{break-inside:avoid}h2.sec{break-after:avoid}'
            '*{-webkit-print-color-adjust:exact;print-color-adjust:exact}}')


def _b64(path):
    return base64.b64encode(Path(path).read_bytes()).decode()


def excel_panel(wb, ctx, charts):
    """Лист «Панель» первым: плитки с цифрами и графики в две колонки."""
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet('Панель', 0)
    ws.sheet_view.showGridLines = False
    ws['A1'] = ctx.get('title') or f'Яндекс Метрика: {ctx["site"]}'
    ws['A1'].font = Font(bold=True, size=16)
    ws['A2'] = period_line(ctx) + f' · {ctx["meta"]}'
    ws['A2'].font = Font(bold=True, size=12)
    if ctx.get('warnings'):
        ws['A3'] = ' '.join(ctx['warnings'])
        ws['A3'].font = Font(color='8A5A00')
    tile_fill = PatternFill('solid', fgColor='F3F3F0')
    col = 1
    for label, value, shown, ch, unit in ctx['kpis'][:6]:
        c1 = ws.cell(row=4, column=col, value=label)
        c2 = ws.cell(row=5, column=col, value=shown)
        txt = '' if ch is None else (f'{ch:+.1f}'.replace('.', ',') + (' п.п.' if unit == 'pp' else '%'))
        good = UP_IS_GOOD.get(label)
        color = '898781' if ch is None or good is None or ch == 0 else (
            '006300' if (ch > 0) == good else 'D03B3B')
        c3 = ws.cell(row=6, column=col, value=f'{txt} {ctx.get("vs", "к прошлому периоду")}' if txt else '')
        c1.font, c2.font, c3.font = Font(color='52514E', size=9), Font(bold=True, size=15), Font(color=color, size=9)
        for c in (c1, c2, c3):
            c.fill = tile_fill
            c.alignment = Alignment(horizontal='left', vertical='center')
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 2)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 2)
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 2)
        ws.column_dimensions[get_column_letter(col + 3)].width = 2   # узкий промежуток между плитками
        col += 4
    ws.row_dimensions[5].height = 26
    # Графики — по два в ряд: левый от колонки A, правый от M (≈ 630 px: промежутки
    # между плитками узкие). Ряд занимает высоту большего из двух графиков
    # (строка Excel — 20 пикселей)
    layout, pair = [], []
    for ch in charts:
        if Path(ch[1]).stem in WIDE:            # широкий график — один в ряду
            if pair:
                layout.append(pair)
                pair = []
            layout.append([ch])
        else:
            pair.append(ch)
            if len(pair) == 2:
                layout.append(pair)
                pair = []
    if pair:
        layout.append(pair)
    row = 8
    for line in layout:
        tallest = 0
        for j, (_, path, _) in enumerate(line):
            img = Image(str(path))
            width = 1220 if Path(path).stem in WIDE else 600
            scale = width / img.width
            img.width, img.height = int(img.width * scale), int(img.height * scale)
            ws.add_image(img, f'{"A" if j == 0 else "M"}{row}')
            tallest = max(tallest, img.height)
        row += tallest // 20 + 2
    return ws


# Переключатель периодов: пересчёт идёт через панель (python run.py panel) —
# у файла, открытого двойным щелчком, сервера нет, и кнопки там выключены
SWITCH_JS = """
(function () {
  const box = document.querySelector('.switch');
  if (!box) return;
  const st = box.querySelector('.sst'), bar = box.querySelector('.sbar i');
  const buttons = box.querySelectorAll('button, input');
  if (!box.dataset.cid) return;
  if (location.protocol === 'file:') {
    buttons.forEach(b => b.disabled = true);
    st.textContent = 'Пересчёт работает в панели: python run.py panel';
    return;
  }
  const key = (m, g) => [box.dataset.cid, box.dataset.kind, m, g].join(':');
  async function poll(m, g) {
    const jobs = await (await fetch('/api/jobs')).json();
    const j = jobs[key(m, g)];
    if (j) {
      bar.style.width = Math.round((j.progress || 0) * 100) + '%';
      st.textContent = j.state === 'queued' ? 'в очереди…' : (j.status || '');
      if (j.state === 'done' && j.link) { st.textContent = 'готово, открываю…'; location.href = j.link; return; }
      if (j.state === 'error') {
        box.classList.remove('busy'); buttons.forEach(b => b.disabled = false);
        st.textContent = 'не получилось: ' + (j.error || ''); return;
      }
    }
    setTimeout(() => poll(m, g), 1500);
  }
  async function run(m, g) {
    if (box.classList.contains('busy')) return;
    box.classList.add('busy');
    buttons.forEach(x => x.disabled = true);
    st.textContent = 'пересчитываю…';
    await fetch('/api/run', {method: 'POST', headers: {'Content-Type': 'application/json', 'X-Panel': '1'},
                             body: JSON.stringify({cid: box.dataset.cid, kind: box.dataset.kind, mode: m, goal: g})});
    poll(m, g);
  }
  box.addEventListener('click', e => {
    const b = e.target.closest('button[data-mode]');
    if (!b) return;
    box.querySelectorAll('button[data-mode]').forEach(x => x.classList.toggle('on', x === b));
    run(b.dataset.mode, box.dataset.goal);
  });
  // главные цели: до пяти галочек, «Пересчитать» — отчёт по ним
  const checks = box.querySelectorAll('.glist input');
  checks.forEach(c => c.addEventListener('change', () => {
    if ([...checks].filter(x => x.checked).length > 5) { c.checked = false; st.textContent = 'не больше 5 целей'; }
  }));
  const apply = box.querySelector('.gapply');
  if (apply) apply.addEventListener('click', () => {
    const ids = [...checks].filter(x => x.checked).map(x => x.value);
    if (!ids.length) { st.textContent = 'выберите хотя бы одну цель'; return; }
    box.querySelector('details').open = false;
    run(box.dataset.mode, ids.join(','));
  });
})();
"""


def html_dashboard(path, ctx, charts):
    """Один самодостаточный HTML: картинки встроены, открывается двойным щелчком."""
    def tile(label, shown, ch, unit):
        good = UP_IS_GOOD.get(label)
        if ch is None:
            delta = ''
        else:
            arrow = '▲' if ch > 0 else ('▼' if ch < 0 else '•')
            cls = 'flat' if good is None or ch == 0 else ('up' if (ch > 0) == good else 'down')
            val = f'{ch:+.1f}'.replace('.', ',') + (' п.п.' if unit == 'pp' else '%')
            delta = f'<div class="delta {cls}">{arrow} {val} <span>{ctx.get("vs", "к прошлому периоду")}</span></div>'
        return f'<div class="tile"><div class="label">{lbl(label)}</div><div class="value">{shown}</div>{delta}</div>'

    def block(kpis, pics):
        tiles = ''.join(tile(label, shown, ch, unit) for label, _, shown, ch, unit in kpis)
        cards = ''.join(
            f'<figure class="card{" wide" if Path(p).stem in WIDE else ""}">'
            f'<img src="data:image/png;base64,{_b64(p)}" alt="{title}"><figcaption>{cap}</figcaption></figure>'
            for title, p, cap in pics)
        return f'<section class="tiles">{tiles}</section><section class="grid">{cards}</section>'

    # SEO-отчёт — двумя разделами: «Поиск» и «Весь трафик»; остальные — одним
    sections = ctx.get('sections')
    body = (''.join(f'<h2 class="sec">{name}</h2>{block(kpis, pics)}' for name, kpis, pics in sections)
            if sections else block(ctx['kpis'], charts))
    import html as _html
    from urllib.parse import quote
    sw = ctx.get('switch')
    pdf = (f'<a class="pdfbtn" href="{quote(ctx["pdf_name"])}" download>Скачать PDF</a>'
           if ctx.get('pdf_name') else '')
    switch = f'<div class="switch">{pdf}</div>' if pdf else ''
    if sw:
        modes = [('yoy', 'Год к году'), ('mom', 'Месяц к месяцу'), ('qoq', 'Квартал к кварталу')]
        goal = str(sw.get('goal') or '')
        chosen = goal.split(',') if goal else []
        options = ctx.get('goal_options') or []
        names = [name for gid, name in options if str(gid) in chosen]
        boxes = ''.join(f'<label><input type="checkbox" value="{gid}"{" checked" if str(gid) in chosen else ""}>'
                        f'{_html.escape(_short(name, 70))}</label>' for gid, name in options)
        title = (_html.escape(_short(names[0], 40)) + (f' и ещё {len(names) - 1}' if len(names) > 1 else '')
                 if names else 'не выбрана')
        pick = (f'<details class="goalsel"><summary>{"Главная цель" if len(names) < 2 else "Главные цели"}: '
                f'<b>{title}</b></summary><div class="glist">{boxes}<div class="gfoot"><span>до 5 целей; '
                f'несколько — достижения и конверсия суммой</span><button type="button" class="gapply">'
                f'Пересчитать</button></div></div></details>' if boxes else '')
        switch = (f'<div class="switch" data-cid="{sw["cid"]}" data-kind="{sw["kind"]}" '
                  f'data-mode="{sw["mode"]}" data-goal="{goal}"><span>Сравнить:</span>'
                  + ''.join(f'<button data-mode="{m}"{" class=on" if m == sw["mode"] else ""}>{label}</button>'
                            for m, label in modes)
                  + pick + pdf + '<div class="sbar"><i></i></div><span class="sst"></span></div>')
    sampled = ctx.get('sampled_note') or ''
    periods = '<span class="vs">против</span>'.join(
        f'<div class="period{" prev" if i else ""}"><span>{name}</span><b>{_span(d1, d2)}</b>'
        f'<i>{_days((d2 - d1).days + 1)}</i></div>' for i, (name, d1, d2) in enumerate(ctx['periods']))
    warn = ''.join(f'<p class="warn">{w}</p>' for w in ctx.get('warnings', []))
    html = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{(ctx.get('title') or 'Метрика: ' + ctx['site']).replace('Яндекс ', '')} · {ctx['period_short']}</title>
<style>
:root{{--surface:#fcfcfb;--page:#f5f5f2;--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--line:#e1e0d9;--good:#006300;--bad:#d03b3b}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--page);color:var(--ink);font:15px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}}
.wrap{{max-width:1240px;margin:0 auto;padding:32px 24px 56px}}
header h1{{font-size:26px;margin:0 0 4px;font-weight:600}}
header p{{margin:0;color:var(--ink2)}}
.periods{{display:flex;flex-wrap:wrap;align-items:center;gap:10px;margin:14px 0 6px}}
.period{{background:var(--surface);border:1px solid var(--line);border-left:4px solid #2a78d6;border-radius:10px;
  padding:8px 14px;display:flex;align-items:baseline;gap:10px}}
.period.prev{{border-left-color:#b9b8b1}}
.period span{{font-size:13px;color:var(--ink2)}} .period b{{font-size:21px;font-weight:600}}
.period i{{font-style:normal;font-size:15px;color:var(--ink2)}}
.vs{{color:var(--muted);font-size:14px}} .meta{{font-size:13px;color:var(--muted)}}
.tiles{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin:24px 0}}
.tile{{background:var(--surface);border:1px solid var(--line);border-radius:12px;padding:14px 16px}}
.tile .label{{color:var(--ink2);font-size:13px}}
.tile .value{{font-size:28px;font-weight:600;margin:2px 0}}
.delta{{font-size:13px;font-weight:600}}
.delta span{{font-weight:400;color:var(--muted)}}
.warn{{margin:12px 0 0;padding:10px 14px;background:#fff6e6;border:1px solid #f1d39a;border-radius:10px;color:#5c4400}}
.delta.up{{color:var(--good)}} .delta.down{{color:var(--bad)}} .delta.flat{{color:var(--muted)}}
.grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(520px,1fr));gap:16px}}
.card{{margin:0;background:var(--surface);border:1px solid var(--line);border-radius:12px;padding:12px 12px 10px}}
.card.wide{{grid-column:1/-1}}
.card img{{width:100%;height:auto;display:block}}
.card figcaption{{color:var(--muted);font-size:12.5px;padding:6px 4px 0}}
h2.sec{{font-size:21px;font-weight:600;margin:34px 0 0;padding-top:18px;border-top:2px solid var(--line)}}
h2.sec:first-child{{border-top:0;margin-top:22px;padding-top:0}}
.switch{{display:flex;flex-wrap:wrap;align-items:center;gap:8px;margin:14px 0 4px}}
.switch>span{{color:var(--ink2);font-size:14px}}
.switch button{{padding:7px 14px;border-radius:8px;border:1px solid var(--line);background:var(--surface);
  color:var(--ink2);font:600 14px system-ui,-apple-system,"Segoe UI",sans-serif;cursor:pointer}}
.switch button:hover{{border-color:#2a78d6;color:#2a78d6}}
.switch button.on{{background:#2a78d6;border-color:#2a78d6;color:#fff}}
.switch button:disabled{{cursor:default;opacity:.6}}
.sbar{{width:180px;height:6px;border-radius:3px;background:var(--line);overflow:hidden;display:none}}
.sbar i{{display:block;height:100%;width:0;background:#2a78d6;transition:width .6s}}
.switch.busy .sbar{{display:block}} .sst{{font-size:13px;color:var(--ink2)}}
.goalsel{{position:relative;margin-left:8px;font-size:14px;color:var(--ink2)}}
.goalsel summary{{cursor:pointer;list-style:none;padding:6px 12px;border:1px solid var(--line);border-radius:8px;
  background:var(--surface);max-width:420px;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}
.goalsel summary::after{{content:' ▾';color:var(--muted)}} .goalsel b{{color:var(--ink);font-weight:600}}
.glist{{position:absolute;z-index:5;top:110%;left:0;width:440px;max-height:380px;overflow:auto;background:#fff;
  border:1px solid var(--line);border-radius:10px;box-shadow:0 8px 24px rgba(0,0,0,.12);padding:8px}}
.glist label{{display:flex;gap:8px;align-items:flex-start;padding:5px 6px;border-radius:6px;color:var(--ink);cursor:pointer}}
.glist label:hover{{background:#eef4fc}} .glist input{{margin-top:3px}}
.gfoot{{position:sticky;bottom:-8px;background:#fff;display:flex;gap:10px;align-items:center;
  justify-content:space-between;padding:8px 6px 4px;border-top:1px solid var(--line);margin-top:6px}}
.gfoot span{{font-size:12px;color:var(--muted)}}
.gapply{{padding:7px 14px;border-radius:8px;border:0;background:#2a78d6;color:#fff;font-weight:600;cursor:pointer}}
.pdfbtn{{margin-left:auto;padding:7px 14px;border-radius:8px;border:1px solid #2a78d6;color:#2a78d6;
  font:600 14px system-ui,-apple-system,"Segoe UI",sans-serif;text-decoration:none}}
.pdfbtn:hover{{background:#eef4fc}}
footer{{margin-top:24px;color:var(--muted);font-size:13px}}
@media (max-width:600px){{.grid{{grid-template-columns:1fr}}}}
{BRAND_CSS}{print_css()}</style></head><body>{brand_header(1240)}<div class="wrap">
<header><h1>{ctx.get('title') or 'Яндекс Метрика: ' + ctx['site']}</h1><div class="periods">{periods}</div><p class="meta">{ctx['meta']}</p>{switch}{warn}</header>
{body}
<footer>Выгружено {ctx['exported']} через API Метрики. Все цифры — в Excel-файле рядом: {ctx['xlsx_name']}.{(' ' + sampled) if sampled else ''}</footer>
</div>{'<script>' + SWITCH_JS + '</script>' if sw else ''}</body></html>"""
    Path(path).write_text(html, encoding='utf-8')
    return path
