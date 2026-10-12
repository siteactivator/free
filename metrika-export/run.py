"""Выгрузка отчётов Яндекс Метрики в Excel через API — без веб-интерфейса.

Когда интерфейс Метрики не открывается или нужна выгрузка в Excel, данные
можно забрать напрямую из API: оно живёт отдельно от интерфейса.

    python run.py counters                         список ваших счётчиков
    python run.py report 12345678                  отчёт за 30 дней
    python run.py report site.ru --days 90         счётчик можно назвать по сайту
    python run.py report 12345678 --from 2026-09-01 --to 2026-09-30
    python run.py raw 12345678 --metrics ym:s:visits --dimensions ym:s:lastsignTrafficSource
    python run.py panel                            все панели в браузере: http://127.0.0.1:8790/

Токен — в файле .env рядом со скриптом (METRIKA_TOKEN=...), как его получить —
README.md. Результаты — в data/ рядом со скриптом.
"""
import argparse
import os
import re
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

# Консоль Windows по умолчанию не в UTF-8: без этого кириллица в выводе
# роняет скрипт, если вывод перенаправлен в файл
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding='utf-8', errors='replace')
    except (AttributeError, ValueError):
        pass

try:
    import requests
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as e:
    sys.exit(f'Нет библиотеки ({e.name}). Выполните в этой папке:\n'
             '  python -m pip install -r requirements.txt')

import charts  # noqa: E402  панель с графиками; seaborn подгружается только при сборке

BASE = Path(__file__).resolve().parent
ENV_PATH = BASE / '.env'
DATA = BASE / 'data'
API = 'https://api-metrika.yandex.net'

# Лимиты API: 30 запросов в секунду с IP, 3 параллельных, 5000 в сутки,
# 200 отчётов за 5 минут на счётчик. Полный отчёт — 20–30 запросов, пауза
# между ними держит нас далеко от любого из лимитов
PAUSE = 0.3
RETRIES = 4              # повторы при сбое сети, 5xx и 429
PAGE = 10000             # строк за запрос (API даёт до 100 000)
TOP = 10000              # строк на лист по умолчанию
# «Запрос слишком сложный» (400) при перегрузке API: сначала повтор с паузой,
# потом выборка — 50 %, затем 10 %. Лист с выборкой помечается в Сводке
COMPLEX_RETRIES = 2
SAMPLED_FALLBACK = ('0.5', '0.1')
GOALS_PER_REQUEST = 10   # 2 метрики на цель; в одном запросе — до 20 метрик

# (метрика, заголовок, вид): int — целое, pct — проценты, float — дробное, dur — время
M_SESSION = [
    ('ym:s:visits', 'Визиты', 'int'),
    ('ym:s:users', 'Посетители', 'int'),
    ('ym:s:pageviews', 'Просмотры', 'int'),
    ('ym:s:bounceRate', 'Отказы, %', 'pct'),
    ('ym:s:pageDepth', 'Глубина просмотра', 'float'),
    ('ym:s:avgVisitDurationSeconds', 'Время на сайте', 'dur'),
    ('ym:s:percentNewVisitors', 'Новые посетители, %', 'pct'),
]
M_GOALS = [
    ('ym:s:sumGoalReachesAny', 'Достижения целей', 'int'),
    ('ym:s:anyGoalConversionRate', 'Конверсия в любую цель, %', 'pct'),
]
M_SHORT = [m for m in M_SESSION if m[0] not in ('ym:s:pageviews', 'ym:s:percentNewVisitors')]

# ключ: (лист, [(группировка, заголовок)], метрики, что это)
REPORTS = {
    'sources':   ('Источники', [('ym:s:lastsignTrafficSource', 'Источник')], M_SESSION + M_GOALS,
                  'откуда пришли: поиск, реклама, ссылки, прямые заходы — по последнему значимому источнику'),
    'engines':   ('Поисковые системы', [('ym:s:lastsignSearchEngineRoot', 'Поисковая система')],
                  M_SESSION + M_GOALS, 'переходы из Яндекса, Google и других поисковиков'),
    'phrases':   ('Поисковые фразы', [('ym:s:lastsignSearchPhrase', 'Поисковая фраза')], M_SHORT + M_GOALS[:1],
                  'по каким запросам пришли; большую часть фраз поисковики скрывают — «(не определено)»'),
    'referrals': ('Сайты-источники', [('ym:s:lastsignReferalSource', 'Сайт')], M_SESSION + M_GOALS,
                  'переходы по ссылкам с других сайтов'),
    'ads':       ('Рекламные системы', [('ym:s:lastsignAdvEngine', 'Рекламная система')], M_SESSION + M_GOALS,
                  'переходы по рекламе: Директ, Google Ads и другие'),
    'social':    ('Соцсети', [('ym:s:lastsignSocialNetwork', 'Соцсеть')], M_SESSION + M_GOALS,
                  'переходы из социальных сетей'),
    'utm':       ('UTM-метки', [('ym:s:lastsignUTMSource', 'utm_source'), ('ym:s:lastsignUTMMedium', 'utm_medium'),
                                ('ym:s:lastsignUTMCampaign', 'utm_campaign')], M_SESSION + M_GOALS,
                  'переходы с UTM-метками'),
    'entry':     ('Страницы входа', [('ym:s:startURL', 'Страница входа')], M_SESSION + M_GOALS,
                  'с каких страниц начинались визиты'),
    'pages':     ('Популярные страницы', [('ym:pv:URL', 'Страница')],
                  [('ym:pv:pageviews', 'Просмотры', 'int'), ('ym:pv:users', 'Посетители', 'int')],
                  'самые просматриваемые страницы'),
    'devices':   ('Устройства', [('ym:s:deviceCategory', 'Устройство')], M_SESSION + M_GOALS,
                  'ПК, смартфоны, планшеты'),
    'os':        ('ОС', [('ym:s:operatingSystemRoot', 'Операционная система')], M_SESSION + M_GOALS,
                  'операционные системы'),
    'browsers':  ('Браузеры', [('ym:s:browser', 'Браузер')], M_SESSION + M_GOALS, 'браузеры'),
    'countries': ('Страны', [('ym:s:regionCountry', 'Страна')], M_SESSION + M_GOALS, 'страны'),
    'cities':    ('Города', [('ym:s:regionCity', 'Город')], M_SESSION + M_GOALS, 'города'),
}
SPECIAL = {'dynamics': 'Динамика', 'goals': 'Цели'}
ORDER = ['dynamics', 'sources', 'engines', 'phrases', 'referrals', 'ads', 'social', 'utm',
         'entry', 'pages', 'goals', 'devices', 'os', 'browsers', 'countries', 'cities']
# Типы целей из API — на русский, как в интерфейсе; неизвестный тип остаётся как есть
GOAL_TYPES = {
    'url': 'Посещение страниц', 'number': 'Количество просмотров', 'step': 'Составная цель',
    'action': 'JavaScript-событие', 'visit_duration': 'Время на сайте', 'phone': 'Клик по номеру телефона',
    'email': 'Клик по email', 'form': 'Отправка формы', 'messenger': 'Переход в мессенджер',
    'file': 'Скачивание файла', 'search': 'Поиск по сайту', 'button': 'Клик по кнопке',
    'social': 'Переход в соцсеть', 'payment_system': 'Платёжная система', 'conditional_call': 'Целевой звонок',
    'call': 'Звонок', 'offline': 'Офлайн-конверсия', 'site': 'Посещение сайта',
    'e_purchase': 'Ecommerce: покупка', 'e_cart': 'Ecommerce: корзина',
    'contact_data': 'Контактные данные', 'contact_data_sent': 'Отправка контактных данных',
    'a_begin_checkout': 'Автоцель: начало оформления заказа', 'a_create_order': 'Автоцель: оформление заказа',
}
GROUPS = {'day': ('ym:s:date', 'День'), 'week': ('ym:s:startOfWeek', 'Неделя (с понедельника)'),
          'month': ('ym:s:startOfMonth', 'Месяц')}


def log(*a):
    print(*a, flush=True)


# ─── токен и запросы ────────────────────────────────────────────────────────
def load_env(path):
    """KEY=VALUE из .env в окружение; уже заданные переменные не трогает."""
    if not path.is_file():
        return
    for line in path.read_text(encoding='utf-8-sig').splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue
        key, _, val = line.partition('=')
        os.environ.setdefault(key.strip(), val.strip().strip('"').strip("'"))


def token():
    load_env(ENV_PATH)
    t = os.environ.get('METRIKA_TOKEN') or os.environ.get('METRICA_TOKEN')
    if not t:
        sys.exit(f'METRIKA_TOKEN не найден. Создайте файл {ENV_PATH}\n'
                 'со строкой METRIKA_TOKEN=ваш_токен (образец — .env.example, README, шаг 3).')
    return t


class TooComplex(Exception):
    """API перегружен или запрос тяжёлый: «Запрос слишком сложный»."""


class QueryError(Exception):
    """Ошибка в самом запросе: неизвестная метрика, нет доступа к счётчику."""


SESSION = requests.Session()


def api_get(path, params=None):
    """GET к API Метрики: повторы при сбоях сети, 5xx и 429; понятные ошибки."""
    SESSION.headers['Authorization'] = f'OAuth {token()}'
    for attempt in range(1, RETRIES + 1):
        time.sleep(PAUSE)
        try:
            r = SESSION.get(API + path, params=params, timeout=120)
        except requests.exceptions.RequestException as e:
            if attempt < RETRIES:
                log(f'   сеть: {type(e).__name__}, повтор через {5 * attempt} с…')
                time.sleep(5 * attempt)
                continue
            sys.exit(f'API Метрики недоступно: {e}\nПроверьте интернет и попробуйте позже.')
        if r.status_code == 200:
            return r.json()
        try:
            msg = '; '.join(e.get('message', '') for e in r.json().get('errors', [])) or r.text[:300]
        except ValueError:
            msg = r.text[:300]
        # Неверный токен Метрика отдаёт как 403 «Invalid oauth_token», не 401
        if r.status_code == 401 or (r.status_code == 403 and 'oauth' in msg.lower()):
            sys.exit(f'{r.status_code}: токен не подошёл — истёк, отозван или скопирован не полностью.\n'
                     'Получите новый (README, шаг 3) и замените его в .env.')
        if r.status_code == 403:
            raise QueryError(f'нет доступа (403): {msg}')
        if r.status_code == 429 or r.status_code >= 500:
            if attempt < RETRIES:
                wait = float(r.headers.get('Retry-After', 0) or 0) or 10 * attempt
                log(f'   API ответило {r.status_code}, жду {wait:.0f} с…')
                time.sleep(wait)
                continue
            raise QueryError(f'API ответило {r.status_code}: {msg}')
        if r.status_code == 400 and 'слишком сложн' in msg.lower():
            raise TooComplex(msg)
        raise QueryError(f'{r.status_code}: {msg}')
    raise QueryError('не удалось получить ответ')


def stat(counter, metrics, dims, d1, d2, top=TOP, sort=None):
    """Отчёт /stat/v1/data со страницами. Слишком сложный — повтор, потом выборка.

    Возвращает строки [(значения группировок, значения метрик)], итоги и пометку
    о выборке: None — данные полные, иначе доля выборки."""
    accuracies = ['full'] * (1 + COMPLEX_RETRIES) + list(SAMPLED_FALLBACK)
    last = None
    for i, acc in enumerate(accuracies):
        try:
            return _stat_pages(counter, metrics, dims, d1, d2, top, sort, acc)
        except TooComplex as e:
            last = e
            if i + 1 < len(accuracies):
                nxt = accuracies[i + 1]
                if nxt == 'full':
                    log(f'   API перегружено («запрос слишком сложный»), повтор через {10 * (i + 1)} с…')
                    time.sleep(10 * (i + 1))
                else:
                    log(f'   всё ещё «слишком сложный» — беру выборку {float(nxt):.0%}')
    raise QueryError(f'API не отдало отчёт даже по выборке 10 %: {last}')


def _stat_pages(counter, metrics, dims, d1, d2, top, sort, accuracy):
    rows, totals, share, total_rows = [], None, 1.0, 0
    offset = 1
    while len(rows) < top:
        params = {'ids': counter, 'metrics': ','.join(metrics), 'date1': d1, 'date2': d2,
                  'limit': min(PAGE, top - len(rows)), 'offset': offset, 'accuracy': accuracy, 'lang': 'ru'}
        if dims:
            params['dimensions'] = ','.join(dims)
        if sort:
            params['sort'] = sort
        j = api_get('/stat/v1/data', params)
        if totals is None:
            totals = j.get('totals') or []
            total_rows = j.get('total_rows', 0)
            if j.get('sampled'):
                share = float(j.get('sample_share') or 1.0)
        data = j.get('data') or []
        for x in data:
            rows.append(([d.get('name') for d in x.get('dimensions', [])], x.get('metrics', [])))
        offset += len(data)
        if not data or offset > total_rows:
            break
    return rows, totals, (None if share >= 1.0 else share), total_rows


# ─── счётчики ───────────────────────────────────────────────────────────────
def all_counters():
    out, offset = [], 1
    while True:
        j = api_get('/management/v1/counters', {'per_page': 1000, 'offset': offset})
        chunk = j.get('counters') or []
        out += chunk
        offset += len(chunk)
        if not chunk or offset > j.get('rows', 0):
            return out


PERMISSION = {'own': 'свой', 'view': 'просмотр', 'edit': 'редактирование'}


def resolve_counter(arg):
    """Номер счётчика или адрес сайта → счётчик из списка доступных."""
    counters = all_counters()
    if arg.isdigit():
        c = next((c for c in counters if str(c['id']) == arg), None)
        if c is None:
            sys.exit(f'Счётчика {arg} нет среди доступных этому токену. Список — python run.py counters')
        return c
    needle = arg.lower().replace('https://', '').replace('http://', '').strip('/')
    def names(c):
        s2 = c.get('site2')
        s2 = s2.get('site', '') if isinstance(s2, dict) else (s2 or '')
        return ' '.join([c.get('site') or '', s2, c.get('name') or '']).lower()
    found = [c for c in counters if needle in names(c)]
    if len(found) > 1:
        # Точное совпадение адреса сайта выигрывает: «site.ru» — это сайт, а не
        # карточки на Картах, у которых в названии тоже есть «site»
        exact = [c for c in found if (c.get('site') or '').lower().strip('/') in (needle, 'www.' + needle)]
        if len(exact) == 1:
            return exact[0]
    if len(found) == 1:
        return found[0]
    if not found:
        sys.exit(f'Не нашёл счётчик по «{arg}». Список — python run.py counters')
    lines = '\n'.join(f'  {c["id"]}  {c.get("site", "")}  {c.get("name", "")}' for c in found[:20])
    sys.exit(f'По «{arg}» подходит несколько счётчиков — укажите номер:\n{lines}')


# ─── Excel ──────────────────────────────────────────────────────────────────
HEAD_FONT = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='4A5D7E')
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill('solid', fgColor='EEF1F6')
FORMATS = {'int': '#,##0', 'pct': '0.0', 'float': '0.00', 'dur': '[h]:mm:ss'}


def cell_value(v, kind):
    if v is None:
        return None
    if kind == 'int':
        return int(round(v))
    if kind == 'dur':
        return v / 86400          # время Excel — доля суток
    return round(v, 2)


def dim_value(v, header):
    if v is None or v == '':
        return '(не определено)'
    if header in ('День', 'Неделя (с понедельника)', 'Месяц') and re.match(r'\d{4}-\d{2}-\d{2}$', str(v)):
        return datetime.strptime(v, '%Y-%m-%d').date()
    return v


def write_table(ws, dim_headers, metrics, rows, totals):
    """Итоги — первой строкой, над шапкой: так они не попадают в автофильтр и не
    уезжают при сортировке."""
    heads = dim_headers + [m[1] for m in metrics]
    if totals:
        ws.append(['Итого и в среднем'] + [None] * (len(dim_headers) - 1)
                  + [cell_value(v, m[2]) for v, m in zip(totals, metrics)])
        for c in ws[1]:
            c.font, c.fill = TOTAL_FONT, TOTAL_FILL
    ws.append(heads)
    hr = ws.max_row
    for c in ws[hr]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[hr].height = 32
    for dims, vals in rows:
        ws.append([dim_value(v, h) for v, h in zip(dims, dim_headers)]
                  + [cell_value(v, m[2]) for v, m in zip(vals, metrics)])
    nd = len(dim_headers)
    for i, m in enumerate(metrics, nd + 1):
        for col in ws.iter_cols(min_col=i, max_col=i):
            for c in col:
                if c.row != hr:
                    c.number_format = FORMATS[m[2]]
    for i, h in enumerate(dim_headers, 1):
        width = 14 if h in ('День', 'Неделя (с понедельника)', 'Месяц') else (60 if h in (
            'Страница входа', 'Страница', 'Поисковая фраза', 'Сайт') else 28)
        ws.column_dimensions[get_column_letter(i)].width = width
        if width == 14:
            for col in ws.iter_cols(min_col=i, max_col=i, min_row=hr + 1):
                for c in col:
                    c.number_format = 'DD.MM.YYYY'
    if totals and nd and ws.column_dimensions['A'].width < 20:
        ws.column_dimensions['A'].width = 20            # чтобы «Итого и в среднем» читалось
    for i in range(nd + 1, nd + len(metrics) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.freeze_panes = ws.cell(row=hr + 1, column=nd + 1)
    ws.auto_filter.ref = f'A{hr}:{get_column_letter(len(heads))}{ws.max_row}'


def note_sheet(ws, text):
    ws['A1'] = text
    ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 100


# ─── отчёт ──────────────────────────────────────────────────────────────────
def period(a):
    if a.date_from or a.date_to:
        if not (a.date_from and a.date_to):
            sys.exit('Нужны обе даты: --from 2026-09-01 --to 2026-09-30')
        d1, d2 = (datetime.strptime(x, '%Y-%m-%d').date() for x in (a.date_from, a.date_to))
    else:
        d2 = date.today() - timedelta(days=1)        # вчера: сегодняшний день неполный
        d1 = d2 - timedelta(days=a.days - 1)
    if d1 > d2:
        sys.exit('Дата начала позже даты конца')
    return d1, d2


def pick_group(a, d1, d2):
    if a.group:
        return a.group
    days = (d2 - d1).days + 1
    return 'day' if days <= 62 else ('week' if days <= 370 else 'month')


def slug(s):
    return re.sub(r'[^\w.-]+', '_', s or '').strip('_')[:60] or 'counter'


def cmd_counters(a):
    counters = all_counters()
    log(f'Доступно счётчиков: {len(counters)}\n')
    log(f'{"Номер":>10}  {"Доступ":<14} Сайт')
    for c in sorted(counters, key=lambda c: (c.get('site') or '')):
        log(f'{c["id"]:>10}  {PERMISSION.get(c.get("permission"), c.get("permission", "")):<14} '
            f'{c.get("site", "")}  {("— " + c["name"]) if c.get("name") else ""}')
    DATA.mkdir(exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Счётчики'
    ws.append(['Номер', 'Сайт', 'Название', 'Доступ', 'Владелец', 'Статус', 'Создан'])
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for c in counters:
        ws.append([c['id'], c.get('site'), c.get('name'), PERMISSION.get(c.get('permission'), c.get('permission')),
                   c.get('owner_login'), c.get('status'), (c.get('create_time') or '')[:10]])
    for col, w in zip('ABCDEFG', (12, 36, 36, 14, 22, 10, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    out = DATA / 'counters.xlsx'
    wb.save(out)
    log(f'\n✅ список сохранён: {out}')


def cmd_report(a):
    counter = resolve_counter(a.counter)
    cid, site = counter['id'], counter.get('site') or counter.get('name') or ''
    d1, d2 = period(a)
    days = (d2 - d1).days + 1
    p2 = d1 - timedelta(days=1)
    p1 = p2 - timedelta(days=days - 1)
    only = [k.strip() for k in a.only.split(',')] if a.only else ORDER
    bad = [k for k in only if k not in REPORTS and k not in SPECIAL]
    if bad:
        sys.exit(f'Неизвестные отчёты: {", ".join(bad)}. Доступны: {", ".join(ORDER)}')
    group = pick_group(a, d1, d2)
    log(f'   счётчик {cid} · {site} · {d1:%d.%m.%Y}–{d2:%d.%m.%Y} ({days} дн.) · '
        f'сравнение с {p1:%d.%m.%Y}–{p2:%d.%m.%Y}')

    want_charts = not a.no_charts
    if want_charts and not charts.available():
        log('   панель с графиками пропущена: нет seaborn — python -m pip install -r requirements.txt')
        want_charts = False

    wb = Workbook()
    summary = wb.active
    summary.title = 'Сводка'
    status = []      # (лист, строк, всего строк, выборка, ошибка)
    tables = {}      # данные листов — для панели с графиками
    goal_list = []   # (цель, достижения)

    # Сводка: итоги периода и прошлого периода
    mtot = M_SESSION + M_GOALS
    try:
        _, cur, s_cur, _ = stat(cid, [m[0] for m in mtot], [], d1.isoformat(), d2.isoformat(), top=1)
        prev = [None] * len(mtot)
        if not a.no_compare:
            _, prev, _, _ = stat(cid, [m[0] for m in mtot], [], p1.isoformat(), p2.isoformat(), top=1)
    except QueryError as e:
        sys.exit(f'Не удалось получить итоги: {e}')

    for key in only:
        if key == 'dynamics':
            dim, head = GROUPS[group]
            title, dims, metrics = SPECIAL[key], [(dim, head)], M_SESSION + M_GOALS
            sort = dim
        elif key == 'goals':
            title = SPECIAL[key]
            status.append(goals_sheet(wb.create_sheet(title), cid, d1, d2, goal_list))
            continue
        else:
            title, dims, metrics, _ = REPORTS[key]
            sort = '-' + metrics[0][0]
        log(f'   {title}…')
        ws = wb.create_sheet(title)
        try:
            rows, totals, sampled, total_rows = stat(cid, [m[0] for m in metrics], [d for d, _ in dims],
                                                     d1.isoformat(), d2.isoformat(), top=a.top, sort=sort)
        except QueryError as e:
            note_sheet(ws, f'Не получено: {e}')
            status.append((title, 0, 0, None, str(e)))
            log(f'   ⚠ {title}: {e}')
            continue
        write_table(ws, [h for _, h in dims], metrics, rows, totals)
        status.append((title, len(rows), total_rows, sampled, None))
        tables[key] = ([h for _, h in dims], metrics, rows, totals)

    write_summary(summary, counter, d1, d2, p1, p2, mtot, cur, prev, s_cur, status, a.no_compare, group)
    folder = DATA / f'{cid}_{slug(site)}'
    folder.mkdir(parents=True, exist_ok=True)
    stem = f'{slug(site)}_{d1.isoformat()}_{d2.isoformat()}'
    out = folder / f'metrika_{stem}.xlsx'

    dashboard = None
    if want_charts:
        # Панель — дополнение: если она не соберётся, Excel всё равно сохранится
        try:
            log('   панель с графиками…')
            ctx = panel_context(counter, site, d1, d2, p1, p2, group, mtot, cur, prev, tables, goal_list,
                                status, out.name, cid, a.no_compare)
            pics = charts.build(ctx, folder / f'charts_{d1.isoformat()}_{d2.isoformat()}')
            charts.excel_panel(wb, ctx, pics)
            dashboard = charts.html_dashboard(folder / f'dashboard_{stem}.html', ctx, pics)
        except Exception as e:  # noqa: BLE001
            log(f'   ⚠ панель с графиками не собрана: {type(e).__name__}: {e}')
            if 'Панель' in wb.sheetnames:
                del wb['Панель']
    wb.save(out)
    got = sum(1 for s in status if not s[4])
    log(f'✅ листов: {got + 1} из {len(status) + 1}' + (f', не получено: {len(status) - got}' if got < len(status) else '')
        + (' + Панель' if dashboard else ''))
    log(f'✅ файл: {out}')
    if dashboard:
        log(f'✅ панель в браузере: {dashboard}')
        log(f'   все панели списком: python run.py panel → http://127.0.0.1:{PORT}/')


def shown(v, kind):
    """Значение для плитки панели: 36 720 / 31,5% / 1,40 / 0:50."""
    if v is None:
        return '—'
    if kind == 'int':
        return f'{v:,.0f}'.replace(',', ' ')
    if kind == 'pct':
        return f'{v:.1f}%'.replace('.', ',')
    if kind == 'dur':
        s = int(round(v))
        return f'{s // 60}:{s % 60:02d}'
    return f'{v:.2f}'.replace('.', ',')


def partial_tail(values, group):
    """1, если последняя точка динамики подозрительно мала — меньше трети медианы.

    Так выглядит день, который Метрика ещё не досчитала (бывает при сбоях и
    в первые часы после полуночи). Для недель и месяцев не проверяем: крайние
    неполные по календарю, это нормально."""
    if group != 'day' or len(values) < 8:
        return 0
    rest = sorted(values[:-1])
    median = rest[len(rest) // 2]
    return 1 if median and values[-1] < median / 3 else 0


def panel_context(counter, site, d1, d2, p1, p2, group, mtot, cur, prev, tables, goal_list, status,
                  xlsx_name, cid, no_compare):
    """Всё, что нужно панели: плитки, динамика текущего и прошлого периода, таблицы."""
    kpis = []
    for (mid, label, kind), v, pv in zip(mtot, cur or [], prev or []):
        if kind == 'pct':
            ch, unit = (round(v - pv, 1) if v is not None and pv is not None else None), 'pp'
        else:
            ch, unit = (round((v - pv) / pv * 100, 1) if v is not None and pv else None), 'pct'
        kpis.append((label, v, shown(v, kind), ch, unit))
    dynamics = None
    if 'dynamics' in tables:
        _, _, rows, _ = tables['dynamics']
        dates = [datetime.strptime(r[0][0], '%Y-%m-%d').date() for r in rows]
        dynamics = {'dates': dates, 'cur': [r[1][0] or 0 for r in rows], 'prev': None,
                    'label': {'day': 'дням', 'week': 'неделям', 'month': 'месяцам'}[group]}
        dynamics['partial_tail'] = partial_tail(dynamics['cur'], group)
        if not no_compare:
            dim = GROUPS[group][0]
            try:
                prev_rows, _, _, _ = stat(cid, ['ym:s:visits'], [dim], p1.isoformat(), p2.isoformat(), sort=dim)
                dynamics['prev'] = [r[1][0] or 0 for r in prev_rows]
            except QueryError as e:
                log(f'   прошлый период для графика динамики не получен: {e}')
    sampled = [s[0] for s in status if s[3]]
    warnings = []
    if dynamics and dynamics.get('partial_tail'):
        last = dynamics['dates'][-1]
        warnings.append(f'Данные за {last:%d.%m.%Y}, похоже, неполные: визитов в разы меньше обычного. '
                        f'Метрика могла ещё не обработать этот день — перевыгрузите отчёт позже.')
        log(f'   ⚠ {warnings[-1]}')
    return {
        'warnings': warnings,
        'site': site, 'kpis': kpis, 'tables': tables, 'dynamics': dynamics,
        'goals': [(g.get('name'), reach) for g, reach in goal_list],
        'period_text': f'{d1:%d.%m.%Y}–{d2:%d.%m.%Y}' + (
            '' if no_compare else f' · сравнение с {p1:%d.%m.%Y}–{p2:%d.%m.%Y}') + f' · счётчик {cid}',
        'period_short': f'{d1:%d.%m}–{d2:%d.%m.%Y}',
        'exported': datetime.now().strftime('%d.%m.%Y %H:%M'),
        'xlsx_name': xlsx_name,
        'sampled_note': ('По выборке: ' + ', '.join(sampled) + '.') if sampled else '',
    }


def goals_sheet(ws, cid, d1, d2, collect=None):
    title = ws.title
    log(f'   {title}…')
    try:
        goals = api_get(f'/management/v1/counter/{cid}/goals').get('goals') or []
    except QueryError as e:
        note_sheet(ws, f'Не получено: {e}')
        return (title, 0, 0, None, str(e))
    if not goals:
        note_sheet(ws, 'Целей в счётчике нет.')
        return (title, 0, 0, None, None)
    results, sampled = [], None
    for i in range(0, len(goals), GOALS_PER_REQUEST):
        chunk = goals[i:i + GOALS_PER_REQUEST]
        metrics = [f'ym:s:goal{g["id"]}{m}' for g in chunk for m in ('reaches', 'conversionRate')]
        try:
            _, totals, s, _ = stat(cid, metrics, [], d1.isoformat(), d2.isoformat(), top=1)
        except QueryError as e:
            note_sheet(ws, f'Не получено: {e}')
            return (title, 0, 0, None, str(e))
        sampled = sampled or s
        for j, g in enumerate(chunk):
            results.append((g, totals[2 * j] if totals else None, totals[2 * j + 1] if totals else None))
    results.sort(key=lambda x: -(x[1] or 0))
    if collect is not None:
        collect.extend((g, reach) for g, reach, _ in results)
    heads = ['Номер цели', 'Цель', 'Тип', 'Достижения', 'Конверсия, %']
    ws.append(heads)
    for c in ws[1]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal='center', wrap_text=True)
    for g, reach, conv in results:
        ws.append([g['id'], g.get('name'), GOAL_TYPES.get(g.get('type'), g.get('type')),
                   cell_value(reach, 'int'), cell_value(conv, 'pct')])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        row[0].number_format, row[1].number_format = FORMATS['int'], FORMATS['pct']
    for col, w in zip('ABCDE', (14, 50, 16, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:E{ws.max_row}'
    return (title, len(results), len(results), sampled, None)


def write_summary(ws, counter, d1, d2, p1, p2, metrics, cur, prev, sampled, status, no_compare, group):
    ws.column_dimensions['A'].width = 34
    for col in 'BCD':
        ws.column_dimensions[col].width = 18
    ws.column_dimensions['E'].width = 60
    big = Font(bold=True, size=14)
    ws['A1'] = f'Яндекс Метрика: {counter.get("site") or counter.get("name")}'
    ws['A1'].font = big
    ws['A2'] = f'Счётчик {counter["id"]}' + (f' · {counter["name"]}' if counter.get('name') else '')
    ws['A3'] = (f'Период: {d1:%d.%m.%Y}–{d2:%d.%m.%Y}'
                + ('' if no_compare else f' · прошлый период: {p1:%d.%m.%Y}–{p2:%d.%m.%Y}'))
    ws['A4'] = f'Выгружено {datetime.now():%d.%m.%Y %H:%M} через API Метрики' + (
        f' · итоги по выборке {sampled:.0%}' if sampled else '')
    ws.append([])
    ws.append(['Показатель', 'Период', 'Прошлый период', 'Изменение'])
    hr = ws.max_row
    for c in ws[hr][:4]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for (mid, label, kind), v, pv in zip(metrics, cur or [], prev or []):
        # Доли в процентах сравниваются разницей в пунктах: отказы 20 % → 22 % —
        # это +2 п.п., а не «+10 %»
        change, fmt = None, '+0.0"%";-0.0"%";0"%"'
        if kind == 'pct':
            fmt = '+0.0" п.п.";-0.0" п.п.";0" п.п."'
            if v is not None and pv is not None:
                change = round(v - pv, 1)
        elif v is not None and pv:
            change = round((v - pv) / pv * 100, 1)
        ws.append([label, cell_value(v, kind), cell_value(pv, kind), change])
        r = ws.max_row
        ws.cell(r, 2).number_format = ws.cell(r, 3).number_format = FORMATS[kind]
        ws.cell(r, 4).number_format = fmt
    ws.append([])
    ws.append(['Лист', 'Строк в файле', 'Всего в Метрике', 'Данные', 'Что это'])
    hr = ws.max_row
    for c in ws[hr][:5]:
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    about = {v[0]: v[3] for v in REPORTS.values()}
    about.update({'Динамика': f'визиты и поведение по {"дням" if group == "day" else "неделям" if group == "week" else "месяцам"}',
                  'Цели': 'достижения и конверсия каждой цели'})
    for title, n, total, s, err in status:
        state = f'не получено: {err}' if err else (f'выборка {s:.0%}' if s else 'полные')
        ws.append([title, n, total, state, about.get(title, '')])
        if err:
            ws.cell(ws.max_row, 4).font = Font(color='C0392B')
    ws.freeze_panes = 'A6'


def cmd_raw(a):
    counter = resolve_counter(a.counter)
    d1, d2 = period(a)
    metrics = [m.strip() for m in a.metrics.split(',') if m.strip()]
    dims = [d.strip() for d in (a.dimensions or '').split(',') if d.strip()]
    log(f'   счётчик {counter["id"]} · {d1:%d.%m.%Y}–{d2:%d.%m.%Y}')
    try:
        rows, totals, sampled, total_rows = stat(counter['id'], metrics, dims, d1.isoformat(), d2.isoformat(),
                                                 top=a.top, sort=a.sort)
    except QueryError as e:
        sys.exit(f'Не получено: {e}')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Отчёт'
    write_table(ws, dims, [(m, m, 'float') for m in metrics], rows, totals)
    folder = DATA / f'{counter["id"]}_{slug(counter.get("site"))}'
    folder.mkdir(parents=True, exist_ok=True)
    out = folder / f'raw_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx'
    wb.save(out)
    log(f'✅ строк: {len(rows)} из {total_rows}' + (f' · выборка {sampled:.0%}' if sampled else ''))
    log(f'✅ файл: {out}')


# ─── панели в браузере ──────────────────────────────────────────────────────
PORT = 8790


def cmd_panel(a):
    """Страница со всеми выгруженными панелями: http://127.0.0.1:8790/.

    Открывается только на этом компьютере. Сервер просто отдаёт файлы из data/
    и в API Метрики не ходит: новый отчёт — командой report, в списке он
    появится после обновления страницы."""
    import html
    import webbrowser
    from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
    from urllib.parse import quote

    def index():
        cards = []
        for p in sorted(root.glob('*/dashboard_*.html'), key=lambda f: f.stat().st_mtime, reverse=True):
            cid, _, site = p.parent.name.partition('_')
            m = re.search(r'_(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})\.html$', p.name)
            per = (f'{date.fromisoformat(m[1]):%d.%m.%Y}–{date.fromisoformat(m[2]):%d.%m.%Y}' if m else '')
            made = datetime.fromtimestamp(p.stat().st_mtime).strftime('%d.%m.%Y %H:%M')
            href = quote(f'{p.parent.name}/{p.name}')
            xlsx = p.with_name('metrika_' + p.name[len('dashboard_'):]).with_suffix('.xlsx')
            excel = (f' · <a href="{quote(f"{p.parent.name}/{xlsx.name}")}">Excel</a>' if xlsx.is_file() else '')
            cards.append(f'<li><a class="site" href="{href}">{html.escape(site)}</a>'
                         f'<span>{per} · счётчик {html.escape(cid)} · выгружено {made}{excel}</span></li>')
        body = (f'<ul>{"".join(cards)}</ul>' if cards else
                '<p class="empty">Отчётов пока нет. Выгрузите первый: '
                '<code>python run.py report &lt;номер счётчика или сайт&gt;</code> — и обновите страницу.</p>')
        return f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Метрика: панели</title>
<style>
body{{margin:0;background:#f5f5f2;color:#0b0b0b;font:15px/1.5 system-ui,-apple-system,"Segoe UI",sans-serif}}
.wrap{{max-width:900px;margin:0 auto;padding:32px 16px 56px}}
h1{{font-size:26px;font-weight:600;margin:0 0 4px}} .sub{{color:#52514e;margin:0 0 24px}}
ul{{list-style:none;margin:0;padding:0;display:grid;gap:10px}}
li{{background:#fcfcfb;border:1px solid #e1e0d9;border-radius:12px;padding:14px 16px}}
li span{{display:block;color:#898781;font-size:13.5px}} a{{color:#2a78d6;text-decoration:none}}
a:hover{{text-decoration:underline}} .site{{font-size:18px;font-weight:600}}
.empty{{background:#fcfcfb;border:1px solid #e1e0d9;border-radius:12px;padding:16px}}
</style></head><body><div class="wrap"><h1>Яндекс Метрика: панели</h1>
<p class="sub">Все выгруженные отчёты, свежие сверху. Новый отчёт: <code>python run.py report …</code>
— и обновите страницу.</p>{body}</div></body></html>"""

    class Handler(SimpleHTTPRequestHandler):
        def __init__(self, *args, **kw):
            super().__init__(*args, directory=str(root), **kw)   # наружу из папки отчётов не отдаёт

        def do_GET(self):
            if self.path.split('?')[0] in ('/', '/index.html'):
                page = index().encode('utf-8')
                self.send_response(200)
                self.send_header('Content-Type', 'text/html; charset=utf-8')
                self.send_header('Content-Length', str(len(page)))
                self.end_headers()
                self.wfile.write(page)
            else:
                super().do_GET()

        def log_message(self, *args):            # без строки в терминале на каждый запрос
            pass

    root = Path(a.data).resolve() if a.data else DATA
    root.mkdir(parents=True, exist_ok=True)
    try:
        srv = ThreadingHTTPServer(('127.0.0.1', a.port), Handler)
    except OSError as e:
        sys.exit(f'Порт {a.port} занят ({e}). Если панель уже запущена — откройте '
                 f'http://127.0.0.1:{a.port}/, иначе укажите другой порт: --port {a.port + 1}')
    url = f'http://127.0.0.1:{a.port}/'
    log(f'✅ панели: {url}  (только на этом компьютере; остановить — Ctrl+C)')
    log(f'   отчёты из {root}')
    if not a.no_browser:
        webbrowser.open(url)
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        log('панель остановлена')


def main():
    ap = argparse.ArgumentParser(description='Отчёты Яндекс Метрики в Excel через API')
    sub = ap.add_subparsers(dest='cmd', required=True)
    sub.add_parser('counters', help='список счётчиков, доступных токену')

    def dates(p):
        p.add_argument('--days', type=int, default=30, help='последние N дней до вчера (по умолчанию 30)')
        p.add_argument('--from', dest='date_from', help='начало периода, ГГГГ-ММ-ДД')
        p.add_argument('--to', dest='date_to', help='конец периода, ГГГГ-ММ-ДД')
        p.add_argument('--top', type=int, default=TOP, help=f'строк на лист, по умолчанию {TOP}')

    r = sub.add_parser('report', help='полный отчёт в Excel')
    r.add_argument('counter', help='номер счётчика или адрес сайта')
    dates(r)
    r.add_argument('--group', choices=list(GROUPS), help='динамика по дням, неделям или месяцам (по умолчанию сам)')
    r.add_argument('--only', help=f'только эти листы через запятую: {",".join(ORDER)}')
    r.add_argument('--no-compare', action='store_true', help='без сравнения с прошлым периодом')
    r.add_argument('--no-charts', action='store_true', help='без панели с графиками')

    w = sub.add_parser('raw', help='свой отчёт: любые метрики и группировки')
    w.add_argument('counter', help='номер счётчика или адрес сайта')
    w.add_argument('--metrics', required=True, help='метрики через запятую, например ym:s:visits,ym:s:users')
    w.add_argument('--dimensions', help='группировки через запятую, например ym:s:lastsignTrafficSource')
    w.add_argument('--sort', help='сортировка, например -ym:s:visits')
    dates(w)

    pn = sub.add_parser('panel', help=f'все панели в браузере: http://127.0.0.1:{PORT}/')
    pn.add_argument('--port', type=int, default=PORT, help=f'порт, по умолчанию {PORT}')
    pn.add_argument('--no-browser', action='store_true', help='не открывать браузер самому')
    pn.add_argument('--data', help='папка с отчётами, по умолчанию data рядом со скриптом')

    a = ap.parse_args()
    if getattr(a, 'days', 1) < 1:
        sys.exit('--days должно быть больше нуля')
    try:
        {'counters': cmd_counters, 'report': cmd_report, 'raw': cmd_raw, 'panel': cmd_panel}[a.cmd](a)
    except QueryError as e:
        sys.exit(f'Ошибка API Метрики: {e}')


if __name__ == '__main__':
    main()
